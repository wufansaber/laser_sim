from pathlib import Path
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from PIL import Image

SRC = Path(os.environ["TASK_DOCX"])
OUT = SRC.with_name(SRC.stem + "_原理数学模型补充版.docx")
doc = Document(SRC)
EQ_DIR = Path(r"D:\work\git_share\laser_sim\tmp\holoor_equations")
EQ_DIR.mkdir(parents=True, exist_ok=True)

def delete_paragraph(p):
    el = p._element
    el.getparent().remove(el)
    p._p = p._element = None

def insert_before(anchor, text="", style=None):
    p = anchor.insert_paragraph_before()
    if style:
        p.style = style
    if text:
        p.add_run(text)
    return p

def set_run(run, name="微软雅黑", size=10.5, color=None, italic=False, bold=False):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    for slot in ("eastAsia", "ascii", "hAnsi"):
        rpr.rFonts.set(qn("w:" + slot), name)
    run.font.size = Pt(size)
    run.italic = italic
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def add_heading(anchor, text, level=2):
    p = insert_before(anchor, style=f"Heading {level}")
    set_run(p.add_run(text), size=13 if level == 2 else 12, color="2E74B5", bold=True)
    p.paragraph_format.keep_with_next = True

def add_body(anchor, text):
    p = insert_before(anchor)
    set_run(p.add_run(text), size=10.5)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15

def add_eq(anchor, display, latex, number):
    p = insert_before(anchor)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    path = EQ_DIR / (number.replace("-", "_") + ".png")
    if not path.exists():
        raise FileNotFoundError(f"Missing rendered equation: {path}")
    with Image.open(path) as im:
        width_pt = min(430, im.width / 220 * 72)
    pic_run = p.add_run()
    pic_run.add_picture(str(path), width=Pt(width_pt))
    for node in pic_run._r.xpath('.//wp:docPr'):
        node.set('title', f'Equation {number}')
        node.set('descr', 'LaTeX equation: ' + latex)
    q = insert_before(anchor)
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after = Pt(5)
    set_run(q.add_run("LaTeX: " + latex), name="Consolas", size=7.6, color="667085")

def add_callout(anchor, title, text, fill="FFF2CC"):
    p = insert_before(anchor)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    p._p.get_or_add_pPr().append(parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), fill)))
    set_run(p.add_run(title + "  "), size=10.2, bold=True)
    set_run(p.add_run(text), size=10.2)

h5 = next(p for p in doc.paragraphs if p.text.strip().startswith("5. 模型原理"))
h6 = next(p for p in doc.paragraphs if p.text.strip().startswith("6. 数学模型"))
h7 = next(p for p in doc.paragraphs if p.text.strip().startswith("7. MATLAB"))
paras = doc.paragraphs
i5 = next(i for i,p in enumerate(paras) if p._p is h5._p)
i6 = next(i for i,p in enumerate(paras) if p._p is h6._p)
i7 = next(i for i,p in enumerate(paras) if p._p is h7._p)
for p in list(paras[i5 + 1:i6]):
    delete_paragraph(p)
paras = doc.paragraphs
i6 = next(i for i,p in enumerate(paras) if p._p is h6._p)
i7 = next(i for i,p in enumerate(paras) if p._p is h7._p)
for p in list(paras[i6 + 1:i7]):
    delete_paragraph(p)

# 第5章：HOLO/OR DOE 原理与方案
add_heading(h6, "5.1 DOE 的相位调制机理")
add_body(h6, "HOLO/OR 的 DOE（Diffractive Optical Element）是透明基底上的微结构相位元件。它不是通过吸收刻出强度图案，而是利用表面浮雕高度产生空间变化的光程差，对入射高斯光束的复振幅相位进行编码；后续自由传播或透镜的傅里叶变换再把相位调制转换为目标强度分布。")
add_eq(h6, r"U_0(x,y)=A_0(x,y)\exp[i\phi_0(x,y)]",
       r"U_0(x,y)=A_0(x,y)\exp\!\left[i\phi_0(x,y)\right]", "5-1")
add_eq(h6, r"t_{\mathrm{DOE}}(x,y)=\exp[i\phi_{\mathrm{DOE}}(x,y)]",
       r"t_{\mathrm{DOE}}(x,y)=\exp\!\left[i\phi_{\mathrm{DOE}}(x,y)\right]", "5-2")
add_eq(h6, r"U_0^{+}(x,y)=U_0(x,y)t_{\mathrm{DOE}}(x,y)",
       r"U_0^{+}(x,y)=U_0(x,y)t_{\mathrm{DOE}}(x,y)", "5-3")
add_body(h6, "若 DOE 浮雕高度为 h(x,y)，基底折射率为 n(λ)，空气折射率近似为 1，则薄元件近似下的相位延迟为：")
add_eq(h6, r"\phi_{\mathrm{DOE}}(x,y)=\frac{2\pi}{\lambda}[n(\lambda)-1]h(x,y)\pmod{2\pi}",
       r"\phi_{\mathrm{DOE}}(x,y)=\frac{2\pi}{\lambda}\left[n(\lambda)-1\right]h(x,y)\pmod{2\pi}", "5-4")
add_body(h6, "HOLO/OR 公开说明其 DOE 使用光刻和等离子刻蚀制造相位台阶。实际产品可采用连续或多级量化相位；量化、刻蚀深度误差、零级光和高阶衍射会影响效率和均匀度，但理想仿真首先按纯相位、近乎无吸收元件处理。")

add_heading(h6, "5.2 单元件焦面型与双元件准直型")
add_body(h6, "HOLO/OR 的标准单元件 Top-Hat Beam Shaper 通常不自带聚焦光焦度。它对入射光施加预设计相位，并配合外部聚焦镜在指定焦面或远场形成平顶强度；该平顶面上的相位一般不是平坦的，因此不能直接用来证明“1.5 mm 近似平行光”。")
add_eq(h6, r"U_f(x_f,y_f)=\frac{e^{ikf}}{i\lambda f}e^{\frac{ik}{2f}(x_f^2+y_f^2)}\mathcal{F}\{U_0^+(x,y)\}",
       r"U_f(x_f,y_f)=\frac{\exp(ikf)}{i\lambda f}\exp\!\left[\frac{ik}{2f}(x_f^2+y_f^2)\right]\mathcal{F}\!\left\{U_0^{+}(x,y)\right\}_{f_x=x_f/(\lambda f),\,f_y=y_f/(\lambda f)}", "5-5")
add_body(h6, "本方案需要 HOLO/OR Collimated Beam Shaper Module 的定制线形版本。依据厂商公开的双元件原理，第一元件把高斯能量重新映射为线形平顶振幅，并在第二元件所在平面形成所需强度；第二元件补偿该平面的残余相位，使模块出口同时满足平顶强度和近似平坦波前。")
add_eq(h6, r"U_1^+=U_{\mathrm{in}}e^{i\phi_1},\qquad U_2^-=\mathcal{P}_L\{U_1^+\}",
       r"U_1^{+}=U_{\mathrm{in}}\exp(i\phi_1),\qquad U_2^{-}=\mathcal{P}_{L}\{U_1^{+}\}", "5-6")
add_eq(h6, r"\phi_2(x,y)=-\arg[U_2^-(x,y)]+\phi_{\mathrm{ref}}(x,y)",
       r"\phi_2(x,y)=-\arg\!\left[U_2^{-}(x,y)\right]+\phi_{\mathrm{ref}}(x,y)", "5-7")
add_eq(h6, r"U_{\mathrm{out}}\approx A_{\mathrm{TH}}\exp[i\phi_{\mathrm{ref}}]",
       r"U_{\mathrm{out}}(x,y)=U_2^{-}(x,y)\exp[i\phi_2(x,y)]\approx A_{\mathrm{TH}}(x,y)\exp[i\phi_{\mathrm{ref}}(x,y)]", "5-8")
add_body(h6, "对理想准直输出，参考相位 φref 可取常数或仅含允许的整体倾角；若存在明显二次相位，则对应有限波前曲率，说明光束仍在会聚或发散。")

add_heading(h6, "5.3 本方案中 HOLO/OR DOE 的作用范围")
add_body(h6, "HOLO/OR 模块只负责 X 向线形平顶强度及出口 X 波前，Y 方向保持原输入特性。模块输入为约 6 mm 的圆形准直光，输出设为约 6 mm 宽的准直线光。后级 CXC1/CXC2 是 X 向无焦缩束组，把 6 mm 缩为 1.5 mm；YBE1/YBE2 和 YFL 只作用于 Y。末片 YFL 在 X 平面必须等效为单位矩阵，不得给 X 增加二次聚焦相位。")
add_body(h6, "厂商标准 TopHatSim.p 只输出焦面归一化强度，可用于验证标准 ST-426-I-Y-A 的强度曲线和 MATLAB 调用链，但不能输出最终定制准直模块的专有相位。本项目采用三层证据：HOLO/OR 强度模型、独立双元件复场对比、最终 HOLO/OR 复振幅/相位文件或样机波前实测。")
add_callout(h6, "模型边界", "下述双元件相位设计是对 HOLO/OR 公开结构的独立物理等效模型，不是对厂商专有微结构的反演。只有厂商针对最终料号提供的复场或样机实测才能关闭厂商相位验收。")

# 第6章：仿真数学模型
add_heading(h7, "6.1 入射光束模型")
add_body(h7, "仿真在 X、Y 可分离的标量旁轴近似下进行。模块输入采用准直高斯光束；D=2w 为强度下降到峰值 exp(-2) 处的全直径。")
add_eq(h7, r"U_{\mathrm{in}}=U_0e^{-(x^2+y^2)/w^2},\quad I_{\mathrm{in}}=I_0e^{-2(x^2+y^2)/w^2}",
       r"U_{\mathrm{in}}(x,y)=U_0\exp\!\left[-\frac{x^2+y^2}{w^2}\right],\qquad I_{\mathrm{in}}=|U_{\mathrm{in}}|^2=I_0\exp\!\left[-\frac{2(x^2+y^2)}{w^2}\right]", "6-1")

add_heading(h7, "6.2 一维高斯到平顶的能量映射")
add_body(h7, "HOLO/OR 的实际优化相位属于厂商知识产权；独立复核模型用能量守恒构造第一元件的初始映射。令输入坐标为 ξ、输出坐标为 x，要求两侧累计功率相等：")
add_eq(h7, r"\frac{\int_{-\infty}^{\xi}I_{\mathrm{in}}(s)ds}{\int_{-\infty}^{+\infty}I_{\mathrm{in}}(s)ds}=\frac{\int_{-D_{\mathrm{TH}}/2}^{T(\xi)}I_{\mathrm{TH}}(u)du}{\int_{-D_{\mathrm{TH}}/2}^{+D_{\mathrm{TH}}/2}I_{\mathrm{TH}}(u)du}",
       r"\frac{\int_{-\infty}^{\xi} I_{\mathrm{in}}(s)\,ds}{\int_{-\infty}^{+\infty} I_{\mathrm{in}}(s)\,ds}=\frac{\int_{-D_{\mathrm{TH}}/2}^{T(\xi)} I_{\mathrm{TH}}(u)\,du}{\int_{-D_{\mathrm{TH}}/2}^{+D_{\mathrm{TH}}/2} I_{\mathrm{TH}}(u)\,du}", "6-2")
add_body(h7, "为避免理想硬边导致的高频振铃，数值对比采用 m=8 的超高斯目标振幅：")
add_eq(h7, r"A_{\mathrm{TH}}(x)=A_0\exp[-|x/(D_{\mathrm{TH}}/2)|^m],\qquad m=8",
       r"A_{\mathrm{TH}}(x)=A_0\exp\!\left[-\left|\frac{x}{D_{\mathrm{TH}}/2}\right|^m\right],\qquad m=8", "6-3")
add_body(h7, "若两相位元件间距为 L，在小角度几何近似下，第一相位的局部梯度使光线从 ξ 偏转到 T(ξ)：")
add_eq(h7, r"\frac{d\phi_1}{d\xi}\approx k\frac{T(\xi)-\xi}{L},\qquad k=\frac{2\pi}{\lambda}",
       r"\frac{d\phi_1}{d\xi}\approx k\,\frac{T(\xi)-\xi}{L},\qquad k=\frac{2\pi}{\lambda}", "6-4")

add_heading(h7, "6.3 角谱传播与第二相位面")
add_body(h7, "复场传播采用角谱法。它同时保留振幅和相位，因而能评价模块后的透镜是否破坏相位展平。")
add_eq(h7, r"U(x,y;z)=\mathcal{F}^{-1}\{\mathcal{F}[U(x,y;0)]H(f_x,f_y;z)\}",
       r"U(x,y;z)=\mathcal{F}^{-1}\!\left\{\mathcal{F}[U(x,y;0)]H(f_x,f_y;z)\right\}", "6-5")
add_eq(h7, r"H=\exp[ikz\sqrt{1-\lambda^2f_x^2-\lambda^2f_y^2}]",
       r"H(f_x,f_y;z)=\exp\!\left[ikz\sqrt{1-\lambda^2f_x^2-\lambda^2f_y^2}\right]", "6-6")
add_body(h7, "第二元件相位取为传播到该面的相位共轭，再叠加允许的参考平面相位，从而令出口平均倾角和曲率接近零。之后同一复场继续传播通过 CXC1/CXC2；YFL 在 X 方向不乘任何薄透镜相位。")

add_heading(h7, "6.4 薄透镜与 X 向无焦缩束")
add_eq(h7, r"t_L(x)=\exp[-ikx^2/(2f_x)]",
       r"t_L(x)=\exp\!\left(-i\frac{kx^2}{2f_x}\right)", "6-7")
add_eq(h7, r"P(d)=\begin{bmatrix}1&d\\0&1\end{bmatrix},\quad L(f)=\begin{bmatrix}1&0\\-1/f&1\end{bmatrix}",
       r"P(d)=\begin{bmatrix}1&d\\0&1\end{bmatrix},\qquad L(f)=\begin{bmatrix}1&0\\-1/f&1\end{bmatrix}", "6-8")
add_eq(h7, r"M_{\mathrm{CXC}}=L(-40)P(120)L(160)=\begin{bmatrix}0.25&120\\0&4\end{bmatrix}",
       r"M_{\mathrm{CXC}}=L(-40)P(120)L(160)=\begin{bmatrix}0.25&120\\0&4\end{bmatrix}", "6-9")
add_body(h7, "矩阵 C=0，说明平行输入的中心/边缘光线输出仍无几何会聚；A=0.25，把 6 mm 缩为 1.5 mm。末级 YFL 在 X 平面为单位矩阵 I，因此：")
add_eq(h7, r"M_{X,\mathrm{total}}=P(150)I_{\mathrm{YFL},X}M_{\mathrm{CXC}},\quad I_{\mathrm{YFL},X}=I",
       r"M_{X,\mathrm{total}}=P(150)I_{\mathrm{YFL},X}M_{\mathrm{CXC}},\qquad I_{\mathrm{YFL},X}=\begin{bmatrix}1&0\\0&1\end{bmatrix}", "6-10")

add_heading(h7, "6.5 Y 向聚焦模型")
add_body(h7, "Y 向采用 M² 修正的高斯束腰关系。目标全 1/e² 宽度 dY=1.0 µm、fY=100 mm、λ=355 nm、M²=1.2 时：")
add_eq(h7, r"d_Y=\frac{4M^2\lambda f_Y}{\pi D_Y},\qquad D_Y=\frac{4M^2\lambda f_Y}{\pi d_Y}=54.240\ \mathrm{mm}",
       r"d_Y=\frac{4M^2\lambda f_Y}{\pi D_Y},\qquad D_Y=\frac{4M^2\lambda f_Y}{\pi d_Y}=54.240\ \mathrm{mm}", "6-11")

add_heading(h7, "6.6 相位、准直与验收指标")
add_body(h7, "在有效 X 光斑内对展开相位作二次拟合。平均倾角来自一次项，曲率来自二次项；相位 RMS 仅去除 piston 和平均倾角，保留二次曲率，以免把会聚或发散误判为准直。")
add_eq(h7, r"\phi_{\mathrm{fit}}=a_2x^2+a_1x+a_0,\qquad \bar{\theta}_x=\frac{\lambda}{2\pi}a_1",
       r"\phi_{\mathrm{fit}}(x)=a_2x^2+a_1x+a_0,\qquad \bar{\theta}_x=\frac{\lambda}{2\pi}a_1", "6-12")
add_eq(h7, r"R_x=\frac{\pi}{\lambda a_2},\qquad \sigma_\phi=\sqrt{\frac{1}{N}\sum[\phi-(b_1x+b_0)]^2}",
       r"R_x=\frac{\pi}{\lambda a_2},\qquad \sigma_\phi=\sqrt{\frac{1}{N}\sum_{n=1}^{N}\left[\phi(x_n)-(b_1x_n+b_0)\right]^2}", "6-13")
add_eq(h7, r"\theta_{\mathrm{edge}}=\frac{D_X/2}{|R_x|}",
       r"\theta_{\mathrm{edge}}=\frac{D_X/2}{|R_x|}", "6-14")
add_body(h7, "本项目暂定判据为 |θ̄x|≤10 µrad、σφ≤0.10 rad、θedge≤50 µrad。有限宽度造成的衍射展宽通过复场随 z 的变化另行报告，不把 λ/D 直接当作几何平行度指标。")

add_heading(h7, "6.7 数值流程与模型假设")
steps = [
    "输入 355 nm、6 mm 准直高斯复场，并乘以第一相位面。",
    "采用角谱法传播到第二元件，建立目标平顶振幅并抵消残余相位。",
    "将模块出口复场传播至 CXC1，依次施加 +160 mm 薄透镜、120 mm 传播、-40 mm 薄透镜及到目标面的传播。",
    "X 传播中不施加 YFL 相位；Y 方向独立用高斯束腰模型验证 1.0 µm 焦斑和 54.24 mm 入射直径。",
    "输出线宽、平均倾角、相位 RMS、曲率和 ABCD 状态；取得 HOLO/OR 官方复场后用相同指标复核。",
]
for text in steps:
    p = insert_before(h7, style="List Number")
    set_run(p.add_run(text), size=10.3)
    p.paragraph_format.space_after = Pt(5)
add_callout(h7, "主要假设", "标量相干传播、薄相位元件、薄透镜、理想口径和无像差。真实 DOE 量化台阶、刻蚀误差、零级光、透镜像差、偏心倾斜、窗口楔角及 355 nm 镀膜误差需由厂商模型、Zemax 黑箱和样机测试补充。", "E8EEF5")

doc.core_properties.subject = (doc.core_properties.subject or "") + "；补充 HOLO/OR DOE 原理与 LaTeX 数学模型"
doc.save(OUT)
print(OUT)
