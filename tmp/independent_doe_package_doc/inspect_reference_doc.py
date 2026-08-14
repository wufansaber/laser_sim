from pathlib import Path
from zipfile import ZipFile
from hashlib import sha256
import sys
from docx import Document
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

path = Path(r"D:\work\git_share\laser_sim\deliverables\HOLOOR_CBSM_LINE_355nm_package\HOLOOR_355nm准直线光技术说明.docx")
doc = Document(path)

print("REFERENCE", path)
print("SHA256", sha256(path.read_bytes()).hexdigest())
print("SECTIONS", len(doc.sections), "PARAGRAPHS", len(doc.paragraphs), "TABLES", len(doc.tables))
for i, section in enumerate(doc.sections, 1):
    print("SECTION", i, section.page_width, section.page_height,
          section.left_margin, section.right_margin,
          section.top_margin, section.bottom_margin,
          section.header_distance, section.footer_distance)

roles = ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"]
for role in roles:
    style = doc.styles[role]
    font = style.font
    pf = style.paragraph_format
    print("STYLE", role, "font=", font.name, "size=", font.size,
          "bold=", font.bold, "color=", getattr(font.color, "rgb", None),
          "before=", pf.space_before, "after=", pf.space_after,
          "line=", pf.line_spacing, "align=", pf.alignment,
          "keep_next=", pf.keep_with_next)

print("BODY_PARAGRAPHS")
for i, p in enumerate(doc.paragraphs):
    text = " ".join(p.text.split())
    if text:
        print(f"P{i:03d}\t{p.style.name}\t{text[:220]}")

print("TABLES")
for ti, table in enumerate(doc.tables):
    print("TABLE", ti, "rows", len(table.rows), "cols", len(table.columns))
    for ri, row in enumerate(table.rows[:4]):
        print("  ROW", ri, " | ".join(" ".join(c.text.split())[:100] for c in row.cells))

for part_name, paragraphs in [("HEADER", doc.sections[0].header.paragraphs),
                              ("FOOTER", doc.sections[0].footer.paragraphs)]:
    print(part_name)
    for p in paragraphs:
        print(p.style.name, p.text)

with ZipFile(path) as zf:
    print("PACKAGE_PARTS")
    for info in zf.infolist():
        print(info.filename, info.file_size, sha256(zf.read(info.filename)).hexdigest())
