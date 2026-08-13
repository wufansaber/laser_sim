from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\work\git_share\laser_sim")
PKG = ROOT / "deliverables" / "HOLOOR_ST426_I_Y_A_355nm_package"
OUT = PKG / "HOLOOR_ST-426-I-Y-A_355nm线聚焦技术方案.docx"
RES = PKG / "03_simulation_results"

BLUE = "2E74B5"
DARK_BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
PALE_YELLOW = "FFF2CC"
PALE_GREEN = "E2F0D9"
PALE_RED = "FCE4D6"
GRAY = "F2F4F7"
DARK_GRAY = "475467"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def table(doc, headers, rows, widths, header_fill=GRAY, font_size=8.7):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for idx, text in enumerate(headers):
        c = t.rows[0].cells[idx]
        c.text = str(text)
        set_cell_shading(c, header_fill)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(font_size)
    set_repeat_table_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].text = str(text)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(font_size)
    set_table_geometry(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_body(doc, text="", bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_numbered_sequence(doc, items):
    """Create a real Word numbered list that restarts at 1."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 1
    num_id = (max(num_ids) + 1) if num_ids else 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{(0xA0B00000 + abstract_id):08X}")
    abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    tmpl = OxmlElement("w:tmpl")
    tmpl.set(qn("w:val"), f"{(0xC0D00000 + abstract_id):08X}")
    abstract.append(tmpl)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    aid = OxmlElement("w:abstractNumId")
    aid.set(qn("w:val"), str(abstract_id))
    num.append(aid)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)

    for text in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        ppr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(nid)
        ppr.append(num_pr)
        p.add_run(text)


def add_callout(doc, title, text, fill=PALE_BLUE):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.cell(0, 0)
    set_cell_shading(c, fill)
    set_cell_margins(c, top=130, start=180, bottom=130, end=180)
    p = c.paragraphs[0]
    p.add_run(title + "  ").bold = True
    p.add_run(text)
    set_table_geometry(t, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def add_paragraph_callout(doc, title, text, fill=PALE_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.add_run(title + "  ").bold = True
    p.add_run(text)
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "808080")
        el.set(qn("w:space"), "5")
        borders.append(el)
    ppr.append(borders)
    return p


def add_equation(doc, equation, note=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(equation)
    r.font.name = "Cambria Math"
    r.font.size = Pt(11)
    r.italic = True
    if note:
        n = doc.add_paragraph(note)
        n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        n.paragraph_format.space_after = Pt(6)
        for rr in n.runs:
            rr.font.size = Pt(8.5)
            rr.font.color.rgb = RGBColor.from_string(DARK_GRAY)


def add_figure(doc, path, caption, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    drawings = p._p.xpath('.//wp:docPr')
    if drawings:
        drawings[-1].set("title", caption.split("。", 1)[0])
        drawings[-1].set("descr", caption)
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    for r in cap.runs:
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    return cap


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)
    paragraph.add_run(" 页")


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.widow_control = True

for style_name in ["List Bullet", "List Bullet 2", "List Number"]:
    st = styles[style_name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(3)

for name, size, color, before, after in [
    ("Title", 24, DARK_BLUE, 0, 12),
    ("Subtitle", 13, DARK_GRAY, 0, 16),
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.text = "355 nm HOLO/OR 线形 Top-Hat 聚焦系统技术方案   |   ST-426-I-Y-A"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(DARK_GRAY)
footer = section.footer
fp = footer.paragraphs[0]
fp.add_run("技术文件 · MATLAB R2019b · ")
add_page_number(fp)
for r in fp.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(DARK_GRAY)

doc.core_properties.title = "355 nm HOLO/OR 线形 Top-Hat 聚焦系统技术方案"
doc.core_properties.subject = "ST-426-I-Y-A 厂商模型与 1.5 mm × 1.0 µm 线聚焦方案"
doc.core_properties.author = "Laser Simulation Project"

# Page 1: memo masthead / cover
mast = doc.add_table(rows=2, cols=2)
mast.style = "Table Grid"
mast.cell(0, 0).merge(mast.cell(0, 1))
set_cell_shading(mast.cell(0, 0), DARK_BLUE)
p = mast.cell(0, 0).paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("TECHNICAL DESIGN MEMO")
r.bold = True
r.font.color.rgb = RGBColor(255, 255, 255)
r.font.size = Pt(10)
mast.cell(1, 0).text = "方案编号\nHOLOOR-ST426-355-100"
mast.cell(1, 1).text = "版本 / 日期\nV1.0 / 2026-08-13"
for c in mast.rows[1].cells:
    set_cell_shading(c, LIGHT_BLUE)
    for p in c.paragraphs:
        for r in p.runs:
            r.font.size = Pt(9)
set_table_geometry(mast, [4680, 4680])

p = doc.add_paragraph(style="Title")
p.add_run("355 nm HOLO/OR 线形 Top-Hat\n聚焦系统技术方案")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p = doc.add_paragraph(style="Subtitle")
p.add_run("基于 ST-426-I-Y-A 与厂商 TopHatSim 模型 · 目标 1.5 mm × 1.0 µm · 工作距离 100 mm")

table(doc, ["项目", "内容"], [
    ["选定路线", "HOLO/OR 标准线形 Angular Top-Hat DOE + 分轴柱面后级"],
    ["厂商型号", "ST-426-I-Y-A（官方条目备注 Optimized 355nm；订货前须书面确认）"],
    ["仿真环境", "MATLAB 9.7.0.1190202 (R2019b)，厂商 TopHatSim.p"],
    ["最终指标", "X 全宽 1.5 mm；Y 全宽 1.0 µm；末级工作距离 100 mm"],
    ["系统长度", "激光出口至目标面约 700 mm"],
], [2300, 7060], font_size=9.2)

add_callout(doc, "方案结论", "采用 ST-426-I-Y-A。HOLO/OR 黑箱模型直接给出 f=100 mm 时约 0.452935 mm 的原生线形 Top-Hat；后级 X 向柱面角倍率 3.311734× 将其映射为 1.5 mm，Y 向柱面扩束至 54.24 mm 并由 fY=100 mm 柱面镜聚焦至 1.0 µm。", PALE_GREEN)
add_callout(doc, "验收边界", "厂商 TopHatSim 只输出焦面强度，不输出复场相位和波前。因此本文件对 1.5 mm × 1.0 µm 强度分布给出计算闭环；“X 向终端平行度”为采购/实验验收项，不能由该黑箱模型宣称已验证。", PALE_YELLOW)

p = doc.add_paragraph()
p.add_run("编制目的：").bold = True
p.add_run("形成可供光机设计、采购询价、MATLAB 复算与实验验收共同使用的单一技术基线。")

page_break(doc)

# Page 2: executive summary and contents
add_heading(doc, "执行摘要", 1)
add_body(doc, "本方案遵循以下约束：激光后的第一个光学元件必须为扩束透镜；仅第一组扩束镜采用 X/Y 等光焦度的圆对称透镜；后续光束尺寸按 X、Y 分轴控制；Y 向工作距离取 100 mm；目标焦面形成 X=1.5 mm 的均匀线宽与 Y=1.0 µm 的高斯窄宽。")
add_body(doc, "HOLO/OR ST-426-I-Y-A 是本方案的核心强度整形元件。厂商模型在 355 nm、6 mm 入射光斑、15 mm 通光口径和 100 mm 理想傅里叶透镜条件下生成线形 Top-Hat。R2019b 实际提取的 exp(-2) 原生线宽为 0.4529349 mm，与产品全角 4.5257 mrad 的小角度计算 0.45257 mm 相符。")
add_body(doc, "原生线宽之后不再对 X 方向进行大幅先扩后缩。X 向只设置一组 3.311734× 角倍率的伽利略柱面中继；Y 向则单独扩束至 54.24 mm，以满足 355 nm、M²=1.2、fY=100 mm 条件下的 1.0 µm 全宽要求。最终采用交叉放置的 X/Y 柱面聚焦镜，而不是新增圆对称透镜。")

add_heading(doc, "内容索引", 2)
table(doc, ["章节", "主题", "用途"], [
    ["1–2", "需求、选型与指标", "定义可采购、可测量的技术基线"],
    ["3", "系统方案与元件表", "给出 0–700 mm 光机布置"],
    ["4", "XZ / YZ 光路图", "核对各轴光焦度和口径"],
    ["5–6", "模型原理与数学模型", "说明 HOLO/OR、傅里叶聚焦和分轴缩放"],
    ["7–8", "MATLAB 方法与结果", "说明厂商黑箱调用、指标提取和输出"],
    ["9–11", "公差、风险、验收", "规定供应商确认与实验验证项目"],
    ["附录", "资料包与参考资料", "保证复算可追溯"],
], [1100, 3120, 5140], font_size=9)

add_callout(doc, "关键工程判断", "若“最终平面 X 向平行光”是指输出复场的波前曲率接近无穷，而不是焦面线形强度，则标准 Angular Top-Hat + 傅里叶透镜路线本身不能仅凭强度曲线完成证明。应要求 HOLO/OR 提供复场/Zemax 黑箱，或改订其可产生准直线光的定制模块。", PALE_RED)

page_break(doc)

# Page 3 requirements and selected model
add_heading(doc, "1. 需求与设计输入", 1)
table(doc, ["类别", "参数", "设计值 / 约束", "定义与备注"], [
    ["光源", "波长 λ", "355 nm", "紫外；采购 DOE 与所有镀膜均按 355 nm 确认"],
    ["光源", "光束质量 M²", "1.2（设计输入）", "用于 Y 向衍射极限束腰估算"],
    ["光源", "激光出口直径", "1.0 mm", "全 1/e² 直径；第一元件即扩束负透镜"],
    ["DOE 入射", "光斑直径", "6.0 mm", "圆形、准直、全 1/e² 直径"],
    ["DOE 入射", "通光口径", "15 mm", "HOLO/OR 模型输入；元件外径另按机械规格确认"],
    ["焦面 X", "Top-Hat 全宽", "1.5 mm", "强度 I=exp(-2)·Imax 处的全宽"],
    ["焦面 X", "中央 90% 均匀度", "计算 99.48%", "定义为中央 90% 区间 Imin/Imax"],
    ["焦面 Y", "高斯全宽", "1.0 µm", "强度 I=exp(-2)·Imax 处的全宽"],
    ["结构", "末级工作距离", "100 mm", "末级交叉柱面聚焦镜至目标面"],
    ["结构", "系统总长", "约 700 mm", "激光出口至目标面"],
    ["结构", "圆对称元件", "仅 SBE1/SBE2", "其后均为分轴柱面元件或无 Y 光焦度 DOE"],
], [1350, 1750, 2300, 3960], font_size=8.4)

add_heading(doc, "2. HOLO/OR 型号选取", 1)
add_body(doc, "选定型号：ST-426-I-Y-A。其属于 HOLO/OR Angular Top-Hat / Stable Top-Hat 线形整形器系列。官方产品条目给出全角 4.5257 mrad、线形输出，并在备注中标注“Optimized 355nm”。本方案以 355 nm 作为唯一设计波长输入厂商 MATLAB 模型。")
table(doc, ["厂商参数", "采用值", "在本方案中的作用"], [
    ["型号", "ST-426-I-Y-A", "选定标准 HOLO/OR 线形 Top-Hat DOE"],
    ["优化波长", "355 nm", "TopHatSim 输入 λ=0.355 µm；订货书面确认"],
    ["输入光束直径", "6 mm", "与第一组 6× 圆对称扩束结果一致"],
    ["有效通光口径", "15 mm", "厂商模型 CA 输入"],
    ["全输出角", "4.5257 mrad", "决定 f=100 mm 时约 0.45257 mm 原生线宽"],
    ["过渡区", "0.14 mrad", "对应边缘过渡；最终像面随 X 角倍率缩放"],
    ["形状", "Line", "X 向一维均匀线形强度"],
], [2100, 2100, 5160], font_size=8.8)
add_callout(doc, "订货警示", "官方网页的型号备注与在线波长字段可能存在显示不一致。PO/合同中必须写明“ST-426-I-Y-A, optimized/design wavelength 355 nm”，并让厂商书面确认型号、基底、双面 AR、输入直径、全角、损伤阈值与验收曲线。", PALE_YELLOW)

# Page 4 system scheme and component list
add_heading(doc, "3. 系统方案描述", 1)
add_heading(doc, "3.1 光路功能链", 2)
flow = doc.add_table(rows=1, cols=5)
flow.style = "Table Grid"
labels = [
    "355 nm 激光\nD≈1 mm",
    "SBE1/SBE2\n圆对称 6× 扩束",
    "ST-426-I-Y-A\nX 向 Top-Hat 相位",
    "Y 扩束 + X 角倍率\n分轴柱面调节",
    "交叉柱面 f=100\n1.5 mm × 1.0 µm",
]
for idx, label in enumerate(labels):
    flow.cell(0, idx).text = label
    set_cell_shading(flow.cell(0, idx), LIGHT_BLUE if idx not in (2, 4) else PALE_GREEN)
    for p in flow.cell(0, idx).paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(8.6)
set_table_geometry(flow, [1872] * 5)

add_heading(doc, "3.2 元件与轴向位置", 2)
table(doc, ["ID", "Z / mm", "光焦度轴", "焦距 / 功能", "设计说明"], [
    ["LASER", "0", "—", "355 nm", "TEM00；出口全 1/e² 直径约 1.0 mm"],
    ["SBE1", "30", "X=Y", "f=-10 mm", "第一光学元件；圆对称负透镜"],
    ["SBE2", "80", "X=Y", "f=+60 mm", "圆对称正透镜；与 SBE1 构成 6× 伽利略扩束"],
    ["DOE", "150", "X 相位", "ST-426-I-Y-A", "6 mm 准直入射；Y 向无折射光焦度"],
    ["YBE1", "180", "Y", "fY=-20 mm", "Y 向负柱面扩束镜"],
    ["YBE2", "340.800", "Y", "fY=+180.800 mm", "与 YBE1 构成 9.040× Y 向扩束"],
    ["CXC1", "420", "X", "fX=+150 mm", "X 向柱面角倍率中继第一片"],
    ["CXC2", "524.707", "X", "fX=-45.293 mm", "X 向柱面角倍率 3.311734×；间距 104.707 mm"],
    ["XFL/YFL", "600", "X + Y", "fX=fY=100 mm", "两片正交柱面聚焦镜；不是圆对称镜"],
    ["TARGET", "700", "—", "工作面", "目标强度 1.5 mm × 1.0 µm"],
], [950, 1050, 1200, 1900, 4260], font_size=8.2)

add_body(doc, "设计意图：X 方向在第一组扩束后保持约 6 mm，不再随 Y 向扩束放大到 54.24 mm；只有 CXC1/CXC2 对 X 的角分布进行一次性缩放。Y 方向则由 YBE1/YBE2 独立放大，以便在 100 mm 工作距离下达到 1.0 µm。这样可以降低 X 方向无必要的大口径需求。")

page_break(doc)

# Page 5 XZ
add_heading(doc, "4. 光路图", 1)
add_heading(doc, "4.1 XZ 平面光路", 2)
add_figure(doc, RES / "vendor_xz_optical_path.png", "图 1  HOLO/OR ST-426-I-Y-A 方案 XZ 平面示意图。实线元件对 X 有作用，虚线位置为仅 Y 向作用的柱面元件。", 6.45)
add_body(doc, "XZ 图的红色包络是用于光机布置的几何示意，不代表 TopHatSim 的复场传播。DOE 后 X 向束宽保持约 6 mm；CXC1/CXC2 将厂商原生角分布放大 3.311734×；末级 XFL 在目标面形成 1.5 mm 线形 Top-Hat。")
add_callout(doc, "关于“平行光”", "目标面处画出的 1.5 mm 是焦面强度线宽，而不是两条几何边缘光线严格平行的证明。厂商模型未输出目标面相位，故 X 向微弧度级平行度只能通过厂商复场模型或实测波前确认。", PALE_YELLOW)

page_break(doc)

# Page 6 YZ
add_heading(doc, "4.2 YZ 平面光路", 2)
add_figure(doc, RES / "vendor_yz_optical_path.png", "图 2  HOLO/OR ST-426-I-Y-A 方案 YZ 平面示意图。SBE1/SBE2 为唯一圆对称扩束组；YBE1/YBE2 与 YFL 均为 Y 向柱面元件。", 6.45)
add_body(doc, "YZ 图显示 6 mm 光束只在 Y 向放大到 54.24 mm。末级 YFL 的焦距为 100 mm，对应 100 mm 工作距离；推荐 Y 向有效通光口径不小于 59.66 mm（在 54.24 mm 全 1/e² 直径基础上留 10% 余量）。X 向 CXC1/CXC2 在本平面无光焦度。")
add_callout(doc, "口径约束", "1.0 µm、355 nm、M²=1.2、fY=100 mm 共同决定 Y 向需要约 54.24 mm 的入射直径。这是衍射与工作距离约束的结果。若不允许约 60 mm 的 Y 向有效口径，则必须在光斑、工作距离、M² 或使用更高 NA 的光学结构之间重新权衡。", PALE_RED)

page_break(doc)

# Page 7 principle
add_heading(doc, "5. 模型原理", 1)
add_heading(doc, "5.1 HOLO/OR Angular Top-Hat 工作机理", 2)
add_body(doc, "输入近高斯准直光经过 DOE 后获得一维空间相位调制。该相位在传播到远场或经过无像差傅里叶透镜后，将输入能量重新分配为具有陡峭边缘和低起伏的平台型线强度。DOE 本身不靠吸收形成轮廓，主要通过相位重分布实现能量整形。")
add_equation(doc, "U_f(x) ∝ ℱ{ U_in(ξ) · exp[j φ_DOE(ξ)] } |_(f_x=x/(λf))", "式中 U_in 为 DOE 面复振幅，φ_DOE 为厂商设计相位，f 为傅里叶透镜焦距。")
add_equation(doc, "I_f(x) = |U_f(x)|²", "TopHatSim.p 给出的曲线对应归一化焦面强度 I_f，而不是复振幅 U_f。")
add_body(doc, "对线形元件，可在工程模型中把 X 向 Top-Hat 与 Y 向高斯聚焦近似分离。HOLO/OR 元件在 X 向完成强度整形，后续柱面光学分别管理 X 的线宽映射和 Y 的窄束腰聚焦。")

add_heading(doc, "5.2 厂商 MATLAB 模型的作用范围", 2)
table(doc, ["内容", "TopHatSim.p 是否提供", "本项目处理方式"], [
    ["输入 λ、光斑、CA、焦距、形状", "是", "按 ST-426-I-Y-A 设计输入"],
    ["焦面归一化强度曲线", "是", "提取最长 Line 对象并计算宽度/均匀度"],
    ["估算焦深与过渡区", "是", "保留厂商控制台输出作为参考"],
    ["DOE 连续/量化相位图", "否", "视为厂商知识产权黑箱"],
    ["目标面复场相位/波前", "否", "不能由当前模型证明 X 向平行度"],
    ["真实透镜像差、装调误差", "否", "需 Zemax/实物实验补充"],
], [2700, 2300, 4360], font_size=8.8)

add_callout(doc, "模型解释", "本方案没有使用独立 NumPy/SciPy 模型替代厂商曲线。厂商结果由用户指定的同一 MATLAB R2019b 直接运行 TopHatSim.p 得到；后级解析计算和指标提取同样在该次 MATLAB 流程内完成。", PALE_GREEN)

page_break(doc)

# Page 8 math model
add_heading(doc, "6. 数学模型与参数计算", 1)
add_heading(doc, "6.1 输入高斯光束", 2)
add_equation(doc, "I(r) = I₀ exp(−2r² / w²),    D = 2w = 6 mm", "D 为 DOE 面全 1/e² 强度直径。")

add_heading(doc, "6.2 厂商全角与原生焦面线宽", 2)
add_equation(doc, "d_native(f) = 2f tan(θ_full/2) ≈ f θ_full", "小角近似适用于 θ_full=4.5257 mrad。")
add_equation(doc, "d_native(100 mm) = 2×100×tan(4.5257×10⁻³/2) = 0.45257 mm")
add_body(doc, "R2019b 从 HOLO/OR 曲线提取的 exp(-2) 宽度为 0.4529349 mm；与目录全角计算的相对差约 0.08%，可视为模型取样和边缘定义差异。")

add_heading(doc, "6.3 X 向线宽缩放", 2)
add_equation(doc, "G_X = d_target / d_native = 1.500000 / 0.4529349 = 3.311734")
add_equation(doc, "G_θ = −f₁/f₂  ⇒  f₂ = −150/3.311734 = −45.2935 mm")
add_body(doc, "选择 CXC1 的 fX=+150 mm，CXC2 的 fX=-45.2935 mm；伽利略无焦间距 L=f₁+f₂=104.7065 mm。X 向只执行这一组角倍率，不随 Y 向 9.04× 扩束。")

add_heading(doc, "6.4 Y 向 1.0 µm 聚焦", 2)
add_equation(doc, "d_Y = 4 M² λ f_Y / (π D_Y)")
add_equation(doc, "D_Y = 4×1.2×0.000355×100 / (π×0.001) = 54.2400 mm")
add_body(doc, "因此 YBE1/YBE2 的扩束倍率为 54.2400/6=9.0400×。采用 10% 口径余量后，末级 Y 向推荐有效通光口径 CA_Y≥59.664 mm。")

add_heading(doc, "6.5 二维强度与评价指标", 2)
add_equation(doc, "I(x,y) = I_X(x) · exp(−8y²/d_Y²)")
add_equation(doc, "Uniformity₉₀ = min(I_X) / max(I_X),   x ∈ 中央 90% 线宽")
add_body(doc, "本文件的 X/Y 宽度都以全 1/e² 强度宽度定义，避免把 FWHM、1/e² 半径和几何线长混为一谈。")

page_break(doc)

# Page 9 MATLAB implementation
add_heading(doc, "7. MATLAB R2019b 厂商模型仿真方法", 1)
add_heading(doc, "7.1 运行环境", 2)
table(doc, ["项目", "记录"], [
    ["MATLAB 可执行文件", r"D:\Program Files\Polyspace\R2019b\bin\matlab.exe"],
    ["版本", "9.7.0.1190202 (R2019b)"],
    ["厂商函数", "TopHatSim.p（调用接口参考 TopHatSim.m）"],
    ["厂商模型输入", "CA=15 mm；beamDia=6 mm；f=100 mm；lambda=0.000355 mm；shape='line'"],
    ["复算主程序", "02_matlab_code/run_HOLOOR_ST426_R2019b.m"],
    ["输出目录", "03_simulation_results/reproduction_output"],
], [2600, 6760], font_size=8.7)

add_heading(doc, "7.2 调用与数据处理流程", 2)
add_numbered_sequence(doc, [
    "检查 version('-release') 必须等于 2019b；否则终止运行，避免版本差异被忽略。",
    "把 `01_vendor_model` 与 `02_matlab_code` 加入路径，构造 nominalStruct 和空的 toleranceStruct。",
    "执行 TopHatSim(nominalStruct, toleranceStruct)，从厂商生成的 Figure 中提取样本数最多的数值曲线。",
    "把横坐标由 µm 换算为 mm，强度归一化，计算 exp(-2) 宽度、FWHM 和中央 90% 均匀度。",
    "根据原生宽度计算 X 向角倍率与柱面焦距；根据高斯束腰公式计算 Y 向光束直径与口径。",
    "导出 CSV、PNG、FIG 与 execution_manifest.txt，形成可追溯运行记录。",
])

add_heading(doc, "7.3 厂商模型核心输入示例", 2)
code = [
    "nominalStruct.CA      = 15;",
    "nominalStruct.beamDia = 6;",
    "nominalStruct.f       = 100;",
    "nominalStruct.lambda  = 355e-6;   % mm",
    "nominalStruct.shape   = 'line';",
    "nominalStruct.THSize  = 0.45257;  % mm at f=100 mm",
    "TopHatSim(nominalStruct, toleranceStruct);",
]
t = doc.add_table(rows=1, cols=1)
t.style = "Table Grid"
c = t.cell(0, 0)
set_cell_shading(c, "F7F7F7")
c.text = "\n".join(code)
for p in c.paragraphs:
    for r in p.runs:
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
set_table_geometry(t, [9360])

add_callout(doc, "可复现性", "资料包中的主程序可独立从任意 MATLAB 当前目录运行；它按相对路径定位厂商模型与输出目录，不依赖本项目根目录的工作路径。", PALE_GREEN)

page_break(doc)

# Page 10 results
add_heading(doc, "8. 仿真结果", 1)
add_heading(doc, "8.1 HOLO/OR 原生焦面曲线", 2)
add_figure(doc, RES / "vendor_native_01.png", "图 3  HOLO/OR TopHatSim.p 在 355 nm、D=6 mm、CA=15 mm、f=100 mm 条件下输出的原生线形 Top-Hat。", 5.45)
table(doc, ["指标", "结果", "说明"], [
    ["原生 X 全 1/e² 宽度", "0.4529349 mm", "由厂商图中的数值曲线提取"],
    ["目录全角估算", "0.45257 mm", "2f tan(θ/2)，θ=4.5257 mrad"],
    ["原生估算焦深", "±905.1 µm", "厂商模型控制台输出"],
    ["最终 X 全 1/e² 宽度", "1.500000 mm", "原生曲线横坐标乘 3.311734"],
    ["最终 X FWHM", "1.474496 mm", "同一归一化曲线计算"],
    ["中央 90% 均匀度", "99.4778%", "Imin/Imax"],
], [2700, 2000, 4660], font_size=8.8)

page_break(doc)

add_heading(doc, "8.2 最终二维焦面强度", 2)
add_figure(doc, RES / "vendor_final_2d.png", "图 4  最终归一化二维焦面强度。X 采用厂商 Top-Hat 曲线缩放，Y 采用 1.0 µm 全 1/e² 高斯模型。", 6.3)
table(doc, ["目标", "设计/计算值", "判定"], [
    ["X 全 1/e² 宽度 1.5 mm", "1.500000 mm", "满足强度宽度"],
    ["Y 全 1/e² 宽度 1.0 µm", "1.000 µm", "解析模型满足；需实测"],
    ["末级工作距离 100 mm", "100.000 mm", "满足"],
    ["中央 90% 均匀度 ≥90%", "99.4778%", "模型满足"],
    ["X 向终端平行度（µrad）", "未由厂商模型给出", "必须追加波前验收"],
], [3900, 2600, 2860], font_size=9)
add_callout(doc, "结果解释", "二维图是可分离强度模型的组合结果，不是包含真实透镜像差、DOE 相位误差、表面质量和装调偏差的全系统复场传播。工程放行须以供应商数据与样机实测闭环。", PALE_YELLOW)

page_break(doc)

# Page 12 tolerances
add_heading(doc, "9. 公差、装调与环境要求", 1)
add_heading(doc, "9.1 HOLO/OR 建议的敏感参数", 2)
add_body(doc, "HOLO/OR 的 Top-Hat 说明强调输入波长、光束直径、准直状态、横向偏心和工作面位置都会改变平台均匀度。以下数值作为初始装调预算，最终以该型号正式数据表为准。")
table(doc, ["误差源", "初始控制值", "工程措施"], [
    ["波长偏差", "建议 ≤±2%", "355 nm 激光谱线通常远小于此值；同时核对温漂"],
    ["DOE 面光斑直径", "建议 ≤±5%", "用 UV 光束分析仪在 DOE 面测全 1/e² 直径"],
    ["XY 偏心", "建议 ≤输入直径的 ±5%", "二维微调架；以焦面平台对称性为准"],
    ["DOE 倾斜", "建议 ≤±5°", "实际按反射抑制和像差进一步收紧"],
    ["输入准直", "必须验证", "DOE 前测波前/发散角；不要只凭几何间距"],
    ["焦面离焦", "小于允许焦深", "Z 向精调；以边缘过渡和 Y 束腰共同优化"],
    ["Y 向口径截断", "CA_Y≥59.66 mm", "核对所有 54.24 mm 光束经过的柱面元件净口径"],
], [1800, 2300, 5260], font_size=8.6)

add_heading(doc, "9.2 装调顺序", 2)
table(doc, ["步骤", "操作"], [
    ["1", "不装 DOE 和末级聚焦镜，先调 SBE1/SBE2，使 355 nm 光束在 DOE 位达到 6.0 mm 且尽可能准直。"],
    ["2", "安装 ST-426-I-Y-A，调整中心与转角，使线形方向与 X 轴一致。"],
    ["3", "安装 YBE1/YBE2，只观察 Y 向光斑直径，调至末级镜前 54.24 mm；确认 X 向仍约 6 mm。"],
    ["4", "安装 CXC1/CXC2，以焦面 X 宽度 1.5 mm 和平台对称性优化两片间距及姿态。"],
    ["5", "安装 XFL/YFL 交叉柱面聚焦组，保持最后光学面到目标面 100 mm，先找 Y 最小束腰，再微调 X 平台。"],
    ["6", "记录焦面强度、线宽、均匀度、波前/发散角、功率效率和温漂，作为验收报告。"],
], [900, 8460], font_size=8.7)

page_break(doc)

# Page 13 risks and procurement
add_heading(doc, "10. 风险与采购确认清单", 1)
table(doc, ["风险", "影响", "关闭条件", "责任方"], [
    ["355 nm 型号字段不一致", "订错设计波长或镀膜", "HOLO/OR 正式报价/图纸明确 ST-426-I-Y-A、355 nm", "采购 + 厂商"],
    ["TopHatSim 无相位", "不能验证 X 向微弧度平行度", "取得复场/Zemax 模型或实测 Shack-Hartmann/剪切干涉结果", "光学 + 厂商"],
    ["Y 向约 60 mm 净口径", "元件体积、价格和像差上升", "完成柱面镜可购性、镀膜均匀性和安装口径审查", "光机 + 采购"],
    ["理想傅里叶透镜假设", "真实边缘、均匀度与焦位偏差", "在目标 NA/口径下完成真实面型光学软件复场分析", "光学"],
    ["输入光束非理想", "平台起伏、边缘不对称", "测量 M²、椭圆率、波前并纳入公差蒙特卡洛", "激光 + 光学"],
    ["UV 损伤/污染", "效率衰减或永久损伤", "厂商给出 LIDT；建立洁净、吹扫和功率爬升流程", "系统 + 厂商"],
], [2150, 2350, 3700, 1160], font_size=8.25)

add_heading(doc, "10.1 询价技术条款建议", 2)
table(doc, ["序号", "询价条款"], [
    ["1", "型号/波长：ST-426-I-Y-A，设计与 AR 镀膜均针对 355 nm。"],
    ["2", "入射条件：近高斯、6.0 mm 全 1/e² 直径、准直；请明确允许的 M²、椭圆率和波前误差。"],
    ["3", "输出：线形 Angular Top-Hat，全角 4.5257 mrad；请给出测量定义、均匀度、过渡区、效率和样品检验曲线。"],
    ["4", "机械：有效通光口径、外径、厚度、基底、楔角、表面质量、平面度、推荐安装方向和定位公差。"],
    ["5", "模型：提供 355 nm 条件下可导入 Zemax/Code V 的黑箱，至少包含目标面复场相位或波前信息。"],
    ["6", "可靠性：355 nm LIDT、平均功率/峰值功率限制、污染控制及保修条件。"],
], [900, 8460], font_size=8.7)

add_callout(doc, "放行原则", "在正式 355 nm 数据表和复场/波前证据未取得前，本方案可用于概念设计、光机预布局和询价，但不应把 X 向微弧度平行度标记为‘已通过厂商模型验证’。", PALE_RED)

page_break(doc)

# Page 14 acceptance
add_heading(doc, "11. 样机验收方案", 1)
table(doc, ["验收项目", "测量方法", "判据/输出", "备注"], [
    ["DOE 面光斑", "UV Beam Profiler", "D=6.0 mm，全 1/e²；记录椭圆率", "先于焦面验收"],
    ["输入准直", "剪切板/波前传感器", "记录曲率、RMS 波前与角度", "以厂商容差为最终判据"],
    ["X 线宽", "高倍率成像 + 标定", "全 1/e² 宽度 1.5 mm", "同时报告 FWHM"],
    ["X 均匀度", "焦面强度曲线", "中央 90% Imin/Imax ≥90%", "建议同时报告 RMS 起伏"],
    ["Y 线宽", "显微物镜/刀口法", "全 1/e² 宽度 1.0 µm", "需去卷积探测器 MTF"],
    ["工作距离", "Z 向位移台", "末级光学面至最佳焦面 100 mm", "明确基准面"],
    ["X 向平行度", "Shack-Hartmann/剪切干涉", "按最终 µrad 规范评定", "当前模型不提供判据值"],
    ["效率", "校准功率计", "DOE 前后及目标 ROI 功率", "分离反射、截断和整形损失"],
    ["稳定性", "功率/位置/温漂记录", "按系统工作时长定义", "UV 热漂移需重点观察"],
], [1700, 2500, 3200, 1960], font_size=8.25)

add_heading(doc, "11.1 设计状态判定", 2)
table(doc, ["项目", "当前状态"], [
    ["HOLO/OR 焦面原生 Top-Hat 曲线", "已在 MATLAB R2019b 厂商模型中复算"],
    ["X=1.5 mm 强度宽度与均匀度", "已由厂商曲线缩放和同次 MATLAB 指标提取得出"],
    ["Y=1.0 µm、fY=100 mm 的口径需求", "已由 M² 修正高斯模型计算"],
    ["XZ/YZ 光机布置", "已形成 0–700 mm 初步布局；需结合实物外形深化"],
    ["真实像差、装调公差、UV 损伤", "待元件选定后做公差分析与样机验证"],
    ["X 向终端波前平行度", "未验证；需厂商复场模型或实测"],
], [3900, 5460], font_size=9)

page_break(doc)

# Appendix package
add_heading(doc, "附录 A：交付资料包结构", 1)
table(doc, ["路径", "内容"], [
    ["01_vendor_model/TopHatSim.p", "HOLO/OR 厂商黑箱 MATLAB 模型"],
    ["01_vendor_model/TopHatSim.m", "厂商可读调用接口与参数说明"],
    ["02_matlab_code/run_HOLOOR_ST426_R2019b.m", "资料包复算主程序"],
    ["02_matlab_code/common_two_solution_config.m", "355 nm、型号、目标与光路参数"],
    ["02_matlab_code/run_vendor_solution.m", "厂商模型调用、曲线提取与后级计算"],
    ["02_matlab_code/profile_metrics_common.m", "宽度与均匀度指标函数"],
    ["02_matlab_code/plot_xz_optical_paths.m", "XZ 光路图生成代码"],
    ["02_matlab_code/plot_yz_optical_paths.m", "YZ 光路图生成代码"],
    ["03_simulation_results/vendor_summary.csv", "厂商方案关键参数与计算结果"],
    ["03_simulation_results/vendor_final_x_profile.csv", "最终 X 归一化强度曲线"],
    ["03_simulation_results/vendor_native.fig/png", "厂商模型原生图形"],
    ["03_simulation_results/vendor_xz/yz_optical_path.png", "XZ/YZ 光路示意图"],
    ["03_simulation_results/xz/yz_optical_layout.csv", "元件 ID、位置、光焦度轴和焦距"],
    ["04_official_references/TH_MATLAB_Tutorial.pdf", "HOLO/OR MATLAB 教程"],
    ["04_official_references/HOLOOR_官方资料链接.md", "官方产品页、应用说明和安装手册链接"],
], [4800, 4560], font_size=8.3)

add_heading(doc, "附录 B：参考资料", 1)
refs = [
    "[1] HOLO/OR, Product Families, https://www.holoor.co.il/product-families/ （访问日期：2026-08-13）。",
    "[2] HOLO/OR, Beam Shaper / Top-Hat product page, https://www.holoor.co.il/product/beam-shaper/ （访问日期：2026-08-13）。",
    "[3] HOLO/OR, Beam Shaper Top-Hat application note, https://www.holoor.co.il/application/beam-shaper-top-hat/ （访问日期：2026-08-13）。",
    "[4] HOLO/OR, Top-Hat Installation Manual v2.2, https://www.holoor.co.il/wp-content/uploads/2020/08/TH_installationManual__v2-2.pdf 。",
    "[5] HOLO/OR, TH MATLAB Tutorial，资料包本地文件：04_official_references/TH_MATLAB_Tutorial.pdf。",
    "[6] 本项目 MATLAB R2019b 执行记录与输出：03_simulation_results。",
]
for ref in refs:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    for r in p.runs:
        r.font.size = Pt(9)

add_paragraph_callout(doc, "文档控制", "本文件记录的是 2026-08-13 的概念/初步设计基线。型号、镀膜、光机外形、末级柱面镜像差与波前验收完成后，应发布 V2.0 详细设计版并冻结 BOM、机械接口与验收限值。", PALE_BLUE)

# Avoid row splitting and mark table captions/heading continuity.
for t in doc.tables:
    for row in t.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)

# Remove trailing empty paragraphs that Word can otherwise push to a blank page.
body = doc._element.body
sect_pr = body.sectPr
node = sect_pr.getprevious()
while node is not None and node.tag == qn("w:p") and not "".join(node.itertext()).strip():
    previous = node.getprevious()
    body.remove(node)
    node = previous

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(OUT)
