from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\work\git_share\laser_sim")
PKG = ROOT / "deliverables" / "HOLOOR_CBSM_LINE_355nm_package"
RES = ROOT / "results" / "holoor_collimated_line_R2019b"
OUT = PKG / "HOLOOR_355nm准直线光技术方案.docx"

BLUE = "2E74B5"; DARK = "1F4E79"; INK = "25364A"
LIGHT = "E8EEF5"; GRAY = "F2F4F7"; GREEN = "E2F0D9"
YELLOW = "FFF2CC"; RED = "FCE4D6"; MUTED = "667085"

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); node = pr.find(qn("w:shd"))
    if node is None: node = OxmlElement("w:shd"); pr.append(node)
    node.set(qn("w:fill"), fill)

def margins(cell, top=80, start=120, bottom=80, end=120):
    pr = cell._tc.get_or_add_tcPr(); node = pr.find(qn("w:tcMar"))
    if node is None: node = OxmlElement("w:tcMar"); pr.append(node)
    for name, value in (("top",top),("start",start),("bottom",bottom),("end",end)):
        item = node.find(qn(f"w:{name}"))
        if item is None: item = OxmlElement(f"w:{name}"); node.append(item)
        item.set(qn("w:w"), str(value)); item.set(qn("w:type"), "dxa")

def geometry(t, widths):
    t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    pr=t._tbl.tblPr
    for tag, value in (("tblW",9360),("tblInd",120)):
        n=pr.find(qn(f"w:{tag}"))
        if n is None: n=OxmlElement(f"w:{tag}"); pr.append(n)
        n.set(qn("w:w"),str(value)); n.set(qn("w:type"),"dxa")
    grid=t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        c=OxmlElement("w:gridCol"); c.set(qn("w:w"),str(width)); grid.append(c)
    for row in t.rows:
        trpr=row._tr.get_or_add_trPr(); cs=OxmlElement("w:cantSplit"); trpr.append(cs)
        for i,cell in enumerate(row.cells):
            pr=cell._tc.get_or_add_tcPr(); w=pr.find(qn("w:tcW"))
            if w is None: w=OxmlElement("w:tcW"); pr.append(w)
            w.set(qn("w:w"),str(widths[i])); w.set(qn("w:type"),"dxa"); margins(cell)

def add_table(doc, headers, rows, widths, fs=8.7):
    t=doc.add_table(rows=1, cols=len(headers)); t.style="Table Grid"
    for i,text in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(text); shade(c,GRAY); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.bold=True; r.font.size=Pt(fs)
    hpr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement("w:tblHeader"); rep.set(qn("w:val"),"true"); hpr.append(rep)
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):
            cells[i].text=str(text); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs: r.font.size=Pt(fs)
    geometry(t,widths); doc.add_paragraph().paragraph_format.space_after=Pt(0); return t

def h(doc,text,level=1):
    p=doc.add_heading(text,level=level); p.paragraph_format.keep_with_next=True; return p

def body(doc,text,bold=None):
    p=doc.add_paragraph()
    if bold and text.startswith(bold): p.add_run(bold).bold=True; p.add_run(text[len(bold):])
    else: p.add_run(text)
    return p

def bullet(doc,text): return doc.add_paragraph(text,style="List Bullet")

def callout(doc,title,text,fill=LIGHT):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; c=t.cell(0,0); shade(c,fill); margins(c,130,180,130,180)
    p=c.paragraphs[0]; p.add_run(title+"  ").bold=True; p.add_run(text); geometry(t,[9360]); doc.add_paragraph().paragraph_format.space_after=Pt(0)

def equation(doc,text,note=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(text); r.font.name="Cambria Math"; r.font.size=Pt(11); r.italic=True
    if note:
        q=doc.add_paragraph(note); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(6)
        for rr in q.runs: rr.font.size=Pt(8.5); rr.font.color.rgb=RGBColor.from_string(MUTED)

def figure(doc,path,caption,width=6.35):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(width))
    for node in p._p.xpath('.//wp:docPr'): node.set("title",caption); node.set("descr",caption)
    q=doc.add_paragraph(caption); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(8)
    for r in q.runs: r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(MUTED)

def page(doc): doc.add_page_break()

def page_number(p):
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; p.add_run("第 ")
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); p._p.append(fld); p.add_run(" 页")

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(0.492)
styles=doc.styles; normal=styles["Normal"]
normal.font.name="Calibri"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"等线"); normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.10; normal.paragraph_format.widow_control=True
for name in ("List Bullet","List Number"):
    st=styles[name]; st.font.name="Calibri"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"等线"); st.font.size=Pt(10.5); st.paragraph_format.space_after=Pt(8); st.paragraph_format.line_spacing=1.167
for name,size,color,before,after in (("Title",24,DARK,0,12),("Subtitle",13,MUTED,0,16),("Heading 1",16,BLUE,16,8),("Heading 2",13,BLUE,12,6),("Heading 3",12,DARK,8,4)):
    st=styles[name]; st.font.name="Calibri"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"等线"); st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
hp=sec.header.paragraphs[0]; hp.text="355 nm HOLO/OR 准直线光整形技术方案  |  相位感知版"; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)
fp=sec.footer.paragraphs[0]; fp.add_run("技术文件 · MATLAB R2019b · "); page_number(fp)
for r in fp.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)
doc.core_properties.title="355 nm HOLO/OR 准直线光整形与Y向聚焦技术方案"
doc.core_properties.subject="X向1.5 mm近似平行光，Y向1.0 µm聚焦，含厂商相位验收接口与自研双元件对比仿真"
doc.core_properties.author="Laser Simulation Project"

# Cover
mast=doc.add_table(rows=2,cols=2); mast.style="Table Grid"; mast.cell(0,0).merge(mast.cell(0,1)); shade(mast.cell(0,0),DARK)
p=mast.cell(0,0).paragraphs[0]; r=p.add_run("TECHNICAL DESIGN MEMO"); r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(10)
mast.cell(1,0).text="方案编号\nHOLOOR-CBSM-355-XCOLL-Y100"; mast.cell(1,1).text="版本 / 日期\nV2.0 / 2026-08-13"
for c in mast.rows[1].cells:
    shade(c,LIGHT)
    for p in c.paragraphs:
        for r in p.runs: r.font.size=Pt(9)
geometry(mast,[4680,4680])
p=doc.add_paragraph(style="Title"); p.add_run("355 nm HOLO/OR 准直线光整形\n与 Y 向聚焦技术方案")
p=doc.add_paragraph(style="Subtitle"); p.add_run("X：1.5 mm 近似平行 Top-Hat · Y：1.0 µm 聚焦 · 工作距离：100 mm")
add_table(doc,["项目","基线"],[
    ["选定厂商系列","HOLO/OR Collimated Beam Shaper Module（定制线形；正式料号待询价确认）"],
    ["工程占位型号","CBSM-LINE-355-6MM-CUSTOM（非 HOLO/OR 正式 PN）"],
    ["第一光学元件","SBE1 圆对称扩束负透镜；SBE1/SBE2 是唯一圆对称扩束组"],
    ["末级元件","YFL：仅 Y 向 fY=100 mm；X 向光焦度为 0"],
    ["仿真环境","MATLAB 9.7.0.1190202 (R2019b)"],
],[2300,7060],9.0)
callout(doc,"设计结论","X 向采用 +160/-40 mm、间距 120 mm 的 4× 无焦缩束组，把 6 mm 准直线光缩为 1.5 mm；末级 YFL 在 X 平面为单位矩阵，所以不会把 X 聚焦。Y 向独立扩束至 54.24 mm 后由 fY=100 mm 柱面镜聚焦至 1.0 µm。",GREEN)
callout(doc,"相位结论","厂商标准 TopHatSim.p 只给强度，不能证明准直波前。本方案把 HOLO/OR 复振幅/相位数据设为最终验收必需输入；缺失时使用独立双元件相位模型做设计对比，但不得冒充厂商专有相位。",YELLOW)
page(doc)

# Summary
h(doc,"执行摘要",1)
body(doc,"本版修正了此前将 z=600 mm 元件写成 XFL/YFL 的错误。正确结构是 YFL：只对 Y 聚焦，X 方向没有光焦度；XZ 图中该元件以虚线表示，X 光束在其前后保持约 1.5 mm 宽。")
body(doc,"由于最终 X 指标是“1.5 mm 近似平行光”，所选 HOLO/OR 产品族从焦面型 ST-426-I-Y-A 调整为定制 Collimated Beam Shaper Module。HOLO/OR 官方资料把准直输出定义为双元件模块：第一元件完成强度重分配，第二元件展平相位。标准 ST-426-I-Y-A 仍保留为原厂 MATLAB 强度参考，不承担最终相位验收。")
add_table(doc,["证据层级","可证明内容","当前状态"],[
    ["HOLO/OR TopHatSim.p","标准焦面型线形 Top-Hat 的归一化强度曲线","已在 MATLAB R2019b 复算；仅作强度参考"],
    ["几何 ABCD","无焦缩束组 C=0、YFL 的 X 光焦度=0、边缘光线角度=0","通过"],
    ["独立双元件复场模型","相位展平后穿过实际后级光学的复振幅传播","设计复核通过；不是厂商数据"],
    ["HOLO/OR 官方复场/相位","最终定制模块的强度、相位、曲率和残余倾角","待厂商提供，因此采购验收未关闭"],
],[2400,4400,2560],8.7)
callout(doc,"物理口径","这里的“近似平行”是波前意义：去除 piston 后的平均倾角、相位 RMS 与波前曲率满足限值。有限宽度产生的衍射展宽另行通过复场传播记录，不把 λ/D 直接等同于几何平行度验收。",RED)
h(doc,"内容索引",2)
add_table(doc,["章节","主题"],[
    ["1-2","需求、HOLO/OR 型号选择与证据边界"],["3-4","系统光路、元件表、XZ/YZ 图"],
    ["5-6","厂商/自研模型原理、相位与数学模型"],["7-8","MATLAB 主程序、结果与状态"],
    ["9-10","公差、采购接口与验收条件"],["附录","交付资料包与参考链接"],
],[1500,7860],9.0)
page(doc)

# Requirements
h(doc,"1. 需求与设计输入",1)
add_table(doc,["类别","参数","设计值 / 约束","定义"],[
    ["光源","波长","355 nm","全部 DOE、透镜与镀膜按 355 nm"],
    ["光源","M²","1.2","用于 Y 向衍射极限计算"],
    ["光源","出口直径","1.0 mm","全 1/e² 直径"],
    ["整形模块输入","圆形准直光斑","6.0 mm","SBE1/SBE2 扩束后；X/Y 相同"],
    ["终端 X","线宽","1.5 mm","全 1/e² 宽度，近似平行光"],
    ["终端 X","平均残余倾角","≤10 µrad（暂定）","去除 piston 后的一阶相位斜率"],
    ["终端 X","相位 RMS","≤0.10 rad（暂定）","去除 piston/平均倾角后，保留曲率"],
    ["终端 X","曲率边缘斜率","≤50 µrad（暂定）","半宽除以拟合波前曲率半径"],
    ["终端 Y","聚焦宽度","1.0 µm","全 1/e² 宽度"],
    ["结构","Y 工作距离","100 mm","YFL 至目标面"],
    ["结构","圆对称元件","仅 SBE1/SBE2","后续均为分轴元件或整形模块"],
],[1200,1800,2400,3960],8.15)
h(doc,"2. HOLO/OR 产品选择",1)
body(doc,"最终选择：HOLO/OR Collimated Beam Shaper Module，定制 355 nm、6 mm 输入、X 向线形平顶输出。官方页面给出该系列可定制 lines 等形状，并说明输出形状可在较长传播范围保持准直；正式尺寸、效率、口径和料号必须由 HOLO/OR 报价确认。")
add_table(doc,["项目","最终路线","参考路线"],[
    ["HOLO/OR 系列","Collimated Beam Shaper Module","Beam Shaper / ST-426-I-Y-A"],
    ["输出物理量","准直线形强度 + 展平波前","焦面/远场线形 Top-Hat 强度"],
    ["元件数","厂商说明为双元件模块","单 DOE + 傅里叶透镜"],
    ["对本项目作用","满足 X 近似平行光的产品架构","保留 TopHatSim.p 作强度曲线参考"],
    ["型号状态","定制 RFQ；正式 PN 待厂商确认","ST-426-I-Y-A 已有本地 MATLAB 黑箱"],
],[2300,3600,3460],8.7)
callout(doc,"不得混用","不能把 ST-426-I-Y-A 的焦面强度曲线当作定制准直模块的波前结果，也不能把独立相位模型标注为 HOLO/OR 原厂相位。",YELLOW)
page(doc)

# Layout
h(doc,"3. 系统方案与元件表",1)
h(doc,"3.1 功能链",2)
add_table(doc,["阶段","作用轴","功能"],[
    ["SBE1/SBE2","X + Y","激光后的第一组圆对称 6× 扩束；输出约 6 mm 准直圆斑"],
    ["CBSM1/CBSM2","X","第一元件强度映射，第二元件相位展平；Y 不变"],
    ["YBE1/YBE2","Y","9.040× Y 向扩束，X 不变"],
    ["CXC1/CXC2","X","+160/-40 mm、间距 120 mm 的 4× 无焦缩束，6 mm→1.5 mm"],
    ["YFL","Y","fY=100 mm；X 光焦度=0"],
    ["TARGET","X + Y","X 为 1.5 mm 近似平行线光；Y 为 1.0 µm 焦斑"],
],[1900,1500,5960],8.8)
h(doc,"3.2 元件位置与光焦度",2)
add_table(doc,["ID","Z / mm","作用轴","焦距 / 功能","说明"],[
    ["LASER","0","-","355 nm","出口全 1/e² 直径约 1.0 mm"],
    ["SBE1","30","X=Y","f=-10 mm","第一光学元件；圆对称负透镜"],
    ["SBE2","80","X=Y","f=+60 mm","与 SBE1 构成唯一圆对称 6× 扩束组"],
    ["CBSM1","150","X","强度映射","HOLO/OR 定制模块第一元件"],
    ["CBSM2","220","X","相位展平","HOLO/OR 定制模块第二元件"],
    ["YBE1","250","Y","fY=-20 mm","Y 向负柱面镜"],
    ["YBE2","410.8","Y","fY=+180.8 mm","Y 向 9.040× 扩束"],
    ["CXC1","430","X","fX=+160 mm","X 无焦缩束第一片"],
    ["CXC2","550","X","fX=-40 mm","与 CXC1 间距 120 mm"],
    ["YFL","600","Y only","fY=+100 mm；fX=∞","最终镜只对 Y 聚焦；X 光焦度=0"],
    ["TARGET","700","-","工作面","X 1.5 mm 近似平行；Y 1.0 µm"],
],[850,900,1200,2200,4210],8.0)
body(doc,"z=600 mm 的元件必须在采购图和装调图中标注为 YFL，而不是 XFL/YFL。其柱轴方向应使 X 平面看到平行平板或零功率通道；不允许在 X 方向引入 f=100 mm 的二次相位。")
page(doc)

# figures
h(doc,"4. XZ / YZ 光路图",1)
h(doc,"4.1 XZ 平面：X 近似平行光",2)
figure(doc,RES/"holoor_collimated_xz_optical_path.png","图 1  最终 HOLO/OR 准直线光方案 XZ 光路。YFL 为虚线，表示 X 光焦度为 0。",6.45)
body(doc,"SBE1/SBE2 将 X 从 1 mm 扩到 6 mm。CBSM1/CBSM2 后 X 保持准直；CXC1/CXC2 只做一次无焦缩束，z=550 mm 后到目标面保持约 1.5 mm。图中的红色边界是近轴几何包络；是否达到微弧度波前指标必须使用复场相位计算或波前实测。")
callout(doc,"关键修正","YFL 不改变 X：在 X 平面 M_YFL,X=I。任何 X 包络在 z=600 mm 后发生汇聚的图都是错误版本。",GREEN)
page(doc)
h(doc,"4.2 YZ 平面：Y 向 100 mm 工作距离聚焦",2)
figure(doc,RES/"holoor_collimated_yz_optical_path.png","图 2  最终方案 YZ 光路。YBE1/YBE2 扩束，YFL 在 100 mm 后把 Y 聚焦到 1.0 µm。",6.45)
body(doc,"由 dY=4M²λfY/(πDY) 得 DY=54.24 mm；留 10% 余量时 YFL 建议有效口径不小于 59.66 mm。该大口径只发生在 Y 方向，X 方向不跟随 Y 扩束，因此无需把所有元件做成约 60 mm 的圆口径。")
callout(doc,"结构权衡","100 mm 工作距离与 1.0 µm、355 nm、M²=1.2 共同决定了 Y 向约 54.24 mm 入射直径。若要缩小 Y 元件口径，必须放宽焦斑、工作距离或 M²/光束定义之一。",RED)
page(doc)

# Principles
h(doc,"5. 模型原理与相位证据",1)
h(doc,"5.1 HOLO/OR 双元件准直整形原理",2)
body(doc,"根据 HOLO/OR 对多元件准直整形器的说明，第一元件改变局部光线方向并完成输入高斯能量到目标线形分布的映射；第二元件消除由第一元件和传播产生的相位斜率，使输出波前重新接近平坦。这个结构与仅在傅里叶焦面形成 Top-Hat 的单 DOE 路线物理目标不同。")
equation(doc,"U_out(x)=A_out(x) exp[j φ_out(x)]","准直验收必须同时包含振幅 A_out 与相位 φ_out；只有 |U|² 不足以判定波前。")
equation(doc,"θ̄_x = (λ/2π) · dφ_fit/dx")
equation(doc,"R_x = π / (λ a₂),     φ_fit(x)=a₂x²+a₁x+a₀")
body(doc,"相位 RMS 的计算只去除 piston 与平均倾角，不去除二次相位；这样残余曲率会真实进入准直判定。曲率同时以半宽处边缘斜率 (D/2)/|R| 报告。")
h(doc,"5.2 三类模型的边界",2)
add_table(doc,["模型","输入 / 输出","能否证明相位","用途"],[
    ["HOLO/OR TopHatSim.p","输入 CA、光斑、焦距、波长；输出焦面强度图","否","ST-426 焦面强度参考"],
    ["独立双元件相位替代模型","设计 φ1/φ2，输出复场 U，并经实际后级传播","能证明自研参考结构","在厂商相位缺失时做设计对比"],
    ["HOLO/OR 定制模块复场","厂商提供 U_exit 或 amplitude+phase","是","最终采购和验收证据"],
],[2300,3300,1800,1960],8.4)
callout(doc,"相位文件规则","主程序优先读取 HOLOOR_CBSM_LINE_355_exit_complex_field.mat。文件必须包含 x_mm 与 U_exit，或 x_mm、amplitude、phase_rad。缺文件时 phaseDataAvailable=0、Vendor acceptance pass=0。",YELLOW)
page(doc)

# math
h(doc,"6. 数学模型",1)
h(doc,"6.1 双元件能量映射",2)
equation(doc,"∫_{-∞}^{x₁} |U_in(ξ)|² dξ = ∫_{-D/2}^{x₂} I_target(η) dη","一维能量守恒确定单调坐标映射 x₂=T(x₁)。")
equation(doc,"dφ₁/dx ≈ k (T(x)-x)/L,     φ₂(x₂) ≈ -arg{P_L[U_in e^{jφ₁}]} ")
body(doc,"独立模型用角谱传播算子 P_L 传播第一元件后的复场，再由第二元件抵消输出相位。当前对比模型采用平滑 8 阶超高斯作为理想准直模块的输出振幅，避免把不可制造的硬边阶跃当成真实面形。")
h(doc,"6.2 X 向无焦缩束",2)
equation(doc,"M_CXC = L(-40) · P(120) · L(+160) = [[0.25, 120], [0, 4]]")
body(doc,"矩阵的 C=0，说明输入平行边缘光线输出仍平行；A=0.25，说明 6 mm 宽度变为 1.5 mm。B=120 mm 不影响零角度输入的几何宽度。")
equation(doc,"M_X,total = P(150) · I_YFL,X · M_CXC,     I_YFL,X=[[1,0],[0,1]]")
body(doc,"因此 z=600 mm 的 YFL 在 X 方向既不增加会聚角，也不改变 X 宽度。")
h(doc,"6.3 Y 向焦斑",2)
equation(doc,"D_Y = 4 M² λ f_Y /(π d_Y) = 54.240 mm")
h(doc,"6.4 波前与衍射的区分",2)
body(doc,"平均相位斜率、二次相位曲率属于波前准直指标；有限口径传播引起的边缘衍射和局部相位起伏属于复场演化。验收不再使用 λ/D=237 µrad 作为“平行度”硬边界，而是直接测/算复场相位，并同时记录传播距离内的线宽变化。")
page(doc)

# MATLAB
h(doc,"7. MATLAB R2019b 实现",1)
add_table(doc,["项目","记录"],[
    ["MATLAB","D:\\Program Files\\Polyspace\\R2019b\\bin\\matlab.exe"],
    ["主程序","02_matlab_code/run_HOLOOR_collimated_line_R2019b.m"],
    ["统一配置","02_matlab_code/common_holoor_collimated_config.m"],
    ["HOLO/OR 相位检查","02_matlab_code/run_holoor_collimated_solution.m"],
    ["独立相位对比","02_matlab_code/simulate_holoor_phase_surrogate.m"],
    ["光路图","02_matlab_code/plot_holoor_collimated_paths.m"],
    ["结果目录","03_simulation_results/collimated_line_R2019b"],
],[2800,6560],8.7)
h(doc,"7.1 运行逻辑",2)
for item in [
    "强制检查 version('-release') 为 2019b。",
    "建立 X 向 CXC 与 YFL 单位矩阵，输出几何验证 CSV。",
    "若厂商相位文件存在，则从有效光斑区提取相位，计算平均倾角、RMS 与曲率。",
    "无论厂商相位是否存在，都运行独立双元件复场对比，穿过同一 CXC 后级。",
    "分别输出 geometricPass、designPass 与 acceptancePass；不得用 designPass 替代厂商验收。",
]: bullet(doc,item)
h(doc,"7.2 主程序调用",2)
t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; c=t.cell(0,0); shade(c,"F7F7F7")
c.text="cd('D:/work/git_share/laser_sim/deliverables/HOLOOR_CBSM_LINE_355nm_package/02_matlab_code');\nresult = run_HOLOOR_collimated_line_R2019b();"
for p in c.paragraphs:
    for r in p.runs: r.font.name="Consolas"; r.font.size=Pt(9)
geometry(t,[9360])
callout(doc,"运行环境备注","本次 R2019b 数值计算成功；MATLAB 对用户偏好目录给出权限/图形初始化警告，但 CSV、MAT、PNG 与执行清单均已生成，数值主流程退出码为 0。",YELLOW)
page(doc)

# Results
h(doc,"8. 仿真结果与判定",1)
h(doc,"8.1 几何光路",2)
add_table(doc,["检查项","结果","判定"],[
    ["CXC 无焦条件 C","0 1/mm","通过"],["最后镜 X 光焦度","0 1/mm","通过"],
    ["CXC 后 X 宽度","1.500000 mm","通过"],["目标面 X 几何宽度","1.500000 mm","通过"],
    ["目标面几何边缘光线角度","0 rad","通过"],
],[4200,2700,2460],9.0)
h(doc,"8.2 独立双元件复场对比",2)
figure(doc,RES/"phase_surrogate"/"phase_surrogate_intensity_angle.png","图 3  独立双元件复场模型经实际 CXC 后级传播的强度与局部相位角。该图不是 HOLO/OR 专有相位数据。",6.25)
add_table(doc,["指标","计算值","设计判据"],[
    ["目标 X 全 1/e² 宽度","1.503906 mm","与 1.5 mm 偏差 0.260%"],
    ["目标平均波前倾角","约 0 µrad","≤10 µrad"],
    ["目标相位 RMS","0.03785 rad","≤0.10 rad"],
    ["目标拟合曲率半径","约 -27.43 m","记录项"],
    ["目标曲率边缘斜率","27.34 µrad","≤50 µrad"],
    ["设计复核","通过","独立模型 + 几何共同通过"],
],[3600,2600,3160],8.8)
callout(doc,"如何理解局部角度曲线","局部相位梯度包含有限宽度传播后的衍射起伏，不能把其 RMS 直接当作平均光轴倾角。验收同时限制平均倾角、去倾斜后的相位 RMS 和曲率边缘斜率。",LIGHT)
page(doc)
h(doc,"8.3 厂商相位验收状态",2)
add_table(doc,["项目","当前值","状态"],[
    ["HOLO/OR 复振幅/相位文件","未收到","phaseDataAvailable=0"],
    ["厂商相位 RMS / 倾角 / 曲率","无法计算","phaseVerified=0"],
    ["几何验证","通过","geometricPass=1"],
    ["独立相位对比","通过","designPass=1"],
    ["最终厂商验收","未关闭","acceptancePass=0"],
],[3600,2700,3060],9.0)
body(doc,"这一区分是本方案的核心：当前结果足以支持光机布局、后级透镜方向与询价参数，但不足以声明“HOLO/OR 定制模块已通过微弧度级波前验收”。一旦取得厂商复场文件，把文件放入 01_vendor_model 并重跑同一主程序即可自动更新状态。")
h(doc,"8.4 ST-426-I-Y-A 强度参考",2)
figure(doc,PKG/"03_simulation_results"/"vendor_native_01.png","图 4  本地 HOLO/OR TopHatSim.p 的 ST-426-I-Y-A 焦面线形强度参考。它不代表最终准直模块的相位。",5.35)
body(doc,"该参考运行使用 355 nm、6 mm 输入光斑、15 mm CA、f=100 mm，在 R2019b 中得到约 0.452935 mm 的原生全 1/e² 线宽。它用于证明厂商黑箱的 MATLAB 调用和强度提取链路，不再用于证明终端 X 准直。")

# Tolerances and procurement
h(doc,"9. 公差、装调与后级相位影响",1)
h(doc,"9.1 后级透镜为什么必须纳入相位验证",2)
body(doc,"HOLO/OR 模块的相位展平只定义在模块出口参考面。其后任何 X 向有光焦度的透镜、错误间距、柱面镜轴向误装或窗口楔角都会重新引入相位。因而主程序把模块出口复场继续传播通过 CXC1/CXC2，并把 YFL 在 X 方向建模为单位相位；只检查模块出口相位是不够的。")
add_table(doc,["误差源","X 向影响","控制方式"],[
    ["CXC 间距偏离 120 mm","C 不再为 0，引入残余会聚/发散","装配后用波前曲率回调间距"],
    ["YFL 柱轴方向错误","在 X 引入 f=100 mm 二次相位","安装图标注 Y only；用柱轴基准面验收"],
    ["CXC 焦距/主面偏差","缩束倍率、曲率与目标宽度变化","用实物 EFL/BFL 和主面位置重算"],
    ["模块出口相位非平坦","后级理想也无法恢复准直","要求 HOLO/OR 复场文件与样机波前"],
    ["窗口/保护片楔角","平均倾角变化","纳入完整光路并测平均倾角"],
],[2500,3000,3860],8.5)
h(doc,"9.2 装调顺序",2)
for item in [
    "先调 SBE1/SBE2，使模块输入为 6.0 mm、圆形、准直光；确认它们是唯一圆对称扩束组。",
    "装入 HOLO/OR 双元件模块，校准 X 线方向、模块间距和出口波前。",
    "装 YBE1/YBE2，只看 Y 直径调至约 54.24 mm，同时确认 X 保持约 6 mm。",
    "装 CXC1/CXC2，以 X 宽度和波前曲率共同调到无焦状态，而非只看 1.5 mm 宽度。",
    "最后装 YFL，确认柱轴只对 Y 聚焦；扫描目标面寻找 Y 最小束腰，同时监测 X 波前不变。",
]: bullet(doc,item)
h(doc,"9.3 初始公差",2)
add_table(doc,["项目","暂定限值","备注"],[
    ["模块输入直径","6.0 mm，建议 ±2%","最终按 HOLO/OR 数据表"],
    ["模块输入波前","由厂商规定","须在模块前实测"],
    ["CXC 间距","120 mm 名义值","装调以输出曲率为准"],
    ["YFL X 光焦度","0；不得有设计功率","残余柱面/面形误差纳入波前"],
    ["YFL 有效口径","Y 向 ≥59.66 mm","X 向可按实际 1.5 mm 光束和机械裕量"],
],[2700,2600,4060],8.7)
page(doc)

# Acceptance
h(doc,"10. 采购接口与验收条件",1)
h(doc,"10.1 HOLO/OR 询价必须项",2)
add_table(doc,["序号","询价/数据要求"],[
    ["1","产品族：Collimated Beam Shaper Module；355 nm；6 mm 圆形准直输入；X 向线形平顶输出；Y 不整形。"],
    ["2","输出 X 名义宽度可为 6 mm，以便后级 4× 无焦缩束至 1.5 mm；或由厂商建议更优尺寸。"],
    ["3","提供正式料号、两元件机械图、元件间距/公差、通光口径、镀膜、效率和 LIDT。"],
    ["4","提供模块出口复振幅 U(x,y) 或 amplitude/phase，含坐标、采样、参考面、单位和归一化定义。"],
    ["5","最好同时提供 Zemax/Code V 黑箱，并允许串接 CXC1/CXC2 与 YFL 后复场传播。"],
    ["6","提供实测输出线宽、均匀度、残余倾角、波前 RMS、曲率半径及规定传播距离内的尺寸变化。"],
],[900,8460],8.6)
h(doc,"10.2 样机验收矩阵",2)
add_table(doc,["验收项目","方法","判据 / 输出"],[
    ["模块输入光斑","UV Beam Profiler","6.0 mm 全 1/e²；记录椭圆率"],
    ["模块输入波前","波前传感器/剪切干涉","按厂商输入公差"],
    ["终端 X 宽度","标定成像","1.5 mm 全 1/e²，建议 ±5%"],
    ["终端 X 平均倾角","相位线性拟合","≤10 µrad（合同前确认）"],
    ["终端 X 相位 RMS","去 piston/tilt 后相位","≤0.10 rad（合同前确认）"],
    ["终端 X 曲率","二次相位拟合","边缘斜率 ≤50 µrad（合同前确认）"],
    ["终端 X 传播稳定性","多 Z 面复场/光斑测量","报告线宽与均匀度随 Z 变化，不用 λ/D 代替"],
    ["终端 Y 宽度","显微/刀口去卷积","1.0 µm 全 1/e²"],
    ["Y 工作距离","Z 位移台","YFL 基准面至最佳焦面 100 mm"],
    ["效率与稳定性","功率计、温漂记录","合同约定"],
],[2200,2600,4560],8.2)
callout(doc,"放行门槛","几何与独立相位对比通过，只代表设计可继续。只有 HOLO/OR 官方复场/相位检查和样机波前实测均通过，acceptancePass 才能关闭。",RED)
page(doc)

# Appendix
h(doc,"附录 A：交付资料包",1)
add_table(doc,["路径","内容"],[
    ["01_vendor_model/TopHatSim.p + .m","HOLO/OR 标准焦面强度模型（参考）"],
    ["01_vendor_model/HOLOOR_CBSM_LINE_355_exit_complex_field_README.md","厂商复场文件格式和验收接口"],
    ["02_matlab_code/run_HOLOOR_collimated_line_R2019b.m","最终主程序"],
    ["02_matlab_code/common_holoor_collimated_config.m","统一光路、型号与验收参数"],
    ["02_matlab_code/run_holoor_collimated_solution.m","ABCD、厂商相位与状态判定"],
    ["02_matlab_code/simulate_holoor_phase_surrogate.m","独立双元件复场对比"],
    ["02_matlab_code/plot_holoor_collimated_paths.m","XZ/YZ 图生成"],
    ["03_simulation_results/collimated_line_R2019b","CSV、MAT、PNG、执行清单"],
    ["04_official_references/HOLOOR_官方资料链接.md","产品系列、FAQ 与原理资料"],
],[4900,4460],8.4)
h(doc,"附录 B：官方资料",1)
refs=[
    "[1] HOLO/OR, Collimated Beam Shaper Module, https://www.holoor.co.il/collimated-beam-shaper/（访问：2026-08-13）。",
    "[2] HOLO/OR, Beam Shaper product page, https://www.holoor.co.il/product/beam-shaper/（访问：2026-08-13）。",
    "[3] HOLO/OR, FAQ, https://www.holoor.co.il/faq/（访问：2026-08-13）。",
    "[4] HOLO/OR, Beam Shaper Advantages / Single vs Multi Element, https://www.holoor.co.il/beam-shaper-advantages/（访问：2026-08-13）。",
    "[5] HOLO/OR, Top Hat Beam case study, https://www.holoor.co.il/top-hat-beam/（访问：2026-08-13）。",
    "[6] HOLO/OR, TH MATLAB Tutorial，本资料包 04_official_references/TH_MATLAB_Tutorial.pdf。",
]
for ref in refs:
    p=doc.add_paragraph(ref); p.paragraph_format.left_indent=Inches(0.15); p.paragraph_format.first_line_indent=Inches(-0.15)
    for r in p.runs: r.font.size=Pt(9)
callout(doc,"文件控制","本文件冻结 V2.0 概念/初步设计基线。取得 HOLO/OR 正式 PN、复场数据和真实透镜处方后，应发布 V3.0 详细设计版并重新运行同一 MATLAB R2019b 主程序。",LIGHT)

OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(str(OUT)); print(OUT)
