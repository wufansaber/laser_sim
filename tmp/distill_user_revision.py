from pathlib import Path
from zipfile import ZipFile
import hashlib, json, os
from docx import Document

src = Path(os.environ['TASK_DOCX'])
doc = Document(src)
parts = {}
with ZipFile(src) as z:
    for info in z.infolist():
        data = z.read(info.filename)
        parts[info.filename] = {"size": len(data), "sha256": hashlib.sha256(data).hexdigest()}

artifact = f"""# artifact.md — 用户修改版技术方案编辑契约

## Reference
- path: `{src}`
- sha256: `{hashlib.sha256(src.read_bytes()).hexdigest()}`
- sections: {len(doc.sections)}
- paragraphs: {len(doc.paragraphs)}
- tables: {len(doc.tables)}
- page count before edit: 通过 Word PDF 渲染核对

## Page system
- US Letter portrait, 8.5 x 11 in
- margins: 1.0 in all sides
- header/footer distance: 0.492 in
- one section; preserve headers, footers, page numbers, page geometry

## Typography and components
- Normal body: existing user document styles and direct formatting are authoritative
- Headings: preserve existing Heading 1 / Heading 2 hierarchy
- Equations: clone current centered equation paragraph treatment; add a visible LaTeX source line in Consolas beneath each display equation
- Tables, figures, captions, callouts: preserve unchanged

## Content flow and edit slots
- Preserve paragraphs 0–28 and all tables/figures before section 5
- Expand only section 5 (model principle and phase evidence) and section 6 (mathematical model)
- Preserve sections 7 onward except pagination shifts caused by inserted content
- Add: DOE phase-delay physics, two-element HOLO/OR architecture, vendor black-box boundary, input dependency, Fourier/ASM propagation, coordinate mapping, DOE height-to-phase relation, ABCD train, Y Gaussian focus, phase metrics, and explicit LaTeX source strings
- Do not alter user wording in requirements, optical layout, XZ/YZ diagrams, results, or references

## Package preservation
- Preserve all package parts and relationships except `word/document.xml` and document metadata timestamps modified by save
- Images, headers, footers, styles, numbering, theme and relationships are preserve-only
- package inventory: `tmp/user_revision_parts.json`

## Fidelity gates
- Original file must remain byte-for-byte unchanged
- Edited copy must retain one section and identical page geometry
- Existing figures/tables/header/footer must render unchanged apart from expected page movement
- Every added formula must have both readable mathematical display and LaTeX source
- No clipping, blank pages, broken tables, or missing Chinese/math glyphs
"""
Path(r"D:\work\git_share\laser_sim\tmp\artifact.md").write_text(artifact, encoding="utf-8")
Path(r"D:\work\git_share\laser_sim\tmp\user_revision_parts.json").write_text(json.dumps(parts,indent=2),encoding="utf-8")
print(src)
