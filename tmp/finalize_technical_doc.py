from pathlib import Path
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


src = Path(os.environ["TASK_DOCX"])
out = Path(os.environ["OUTPUT_DOCX"])
eq_dir = Path(os.environ["EQ_DIR"])
eq_dir.mkdir(parents=True, exist_ok=True)
doc = Document(src)


def set_run(run, font="Microsoft YaHei", size=10.5, bold=False,
            color=None, italic=False):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    for slot in ("eastAsia", "ascii", "hAnsi"):
        rpr.rFonts.set(qn("w:" + slot), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph(p):
    element = p._element
    element.getparent().remove(element)
    p._p = p._element = None


def insert_before(anchor, text="", style=None):
    p = anchor.insert_paragraph_before()
    if style:
        p.style = style
    if text:
        set_run(p.add_run(text))
    return p


def add_heading(anchor, text, level=2):
    p = insert_before(anchor, style=f"Heading {level}")
    set_run(p.add_run(text), size=13 if level == 2 else 11.5,
            bold=True, color="2E74B5")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(anchor, text):
    p = insert_before(anchor)
    set_run(p.add_run(text))
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    return p


def add_code(anchor, lines):
    p = insert_before(anchor)
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p._p.get_or_add_pPr().append(
        parse_xml(r'<w:shd {} w:fill="F3F4F6"/>'.format(nsdecls("w"))))
    set_run(p.add_run(lines), font="Consolas", size=8.3, color="202124")
    return p


def add_callout(anchor, title, text, fill="FFF2CC"):
    p = insert_before(anchor)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p._p.get_or_add_pPr().append(
        parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), fill)))
    set_run(p.add_run(title + "  "), bold=True, size=10.2)
    set_run(p.add_run(text), size=10.2)
    return p


def render_eq(key, latex):
    path = eq_dir / f"{key.replace('-', '_')}_final.png"
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def replace_equation_image(latex_prefix, key, latex):
    latex_p = next(p for p in doc.paragraphs if p.text.startswith(latex_prefix))
    idx = next(i for i, p in enumerate(doc.paragraphs)
               if p._p is latex_p._p)
    old_pic = doc.paragraphs[idx - 1]
    remove_paragraph(old_pic)
    p = latex_p.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    pic = p.add_run()
    pic.add_picture(str(render_eq(key, latex)), width=Inches(5.9))
    for node in pic._r.xpath('.//wp:docPr'):
        node.set("title", f"Equation {key}")
        node.set("descr", "LaTeX equation: " + latex)


# Correct a cover typo.
for p in doc.paragraphs:
    if "最有一片镜子" in p.text:
        for r in p.runs:
            if "最有一片镜子" in r.text:
                r.text = r.text.replace("最有一片镜子", "最后一片镜子")

# Remove the obsolete single-element focal-plane discussion.
old_heading = next(p for p in doc.paragraphs
                   if p.text.strip() == "5.2 单元件焦面型与双元件准直型")
old_body = next(p for p in doc.paragraphs
                if p.text.startswith("HOLO/OR 的标准单元件 Top-Hat Beam Shaper"))
old_latex = next(p for p in doc.paragraphs if p.text.startswith("LaTeX: U_f"))
old_latex_idx = next(i for i, p in enumerate(doc.paragraphs)
                     if p._p is old_latex._p)
remove_paragraph(doc.paragraphs[old_latex_idx - 1])
remove_paragraph(old_latex)
old_heading.text = ""
set_run(old_heading.add_run("5.2 HOLO/OR 双元件准直整形方案"),
        size=13, bold=True, color="2E74B5")
remove_paragraph(old_body)

# Close the equation-number gap created by deleting the focal-plane model.
replace_equation_image("LaTeX: U_1^{+}", "5-5",
                       r"U_1^{+}=U_{\mathrm{in}}e^{i\phi_1},\qquad U_2^{-}=\mathcal{P}_{L}\!\left\{U_1^{+}\right\}")
replace_equation_image("LaTeX: \\phi_2", "5-6",
                       r"\phi_2(x,y)=-\arg\!\left[U_2^{-}(x,y)\right]+\phi_{\mathrm{ref}}(x,y)")
replace_equation_image("LaTeX: U_{\\mathrm{out}}", "5-7",
                       r"U_{\mathrm{out}}\approx A_{\mathrm{TH}}e^{i\phi_{\mathrm{ref}}}")

# Remove the old vendor focal-model wording from the principle section.
old_vendor_p = next(p for p in doc.paragraphs
                    if p.text.startswith("厂商标准 TopHatSim.p"))
old_vendor_p.text = ""
set_run(old_vendor_p.add_run(
    "HOLO/OR 未公开最终定制准直模块的专有微结构与出口相位数据。"
    "本资料包因此只保留厂商出口复场数据接口和独立双元件复场替代模型："
    "前者用于收到正式数据后的验收，后者用于在缺少厂商相位时检查后级透镜对整形与波前的影响；"
    "独立模型不得标记为 HOLO/OR 专有模型。"))

# Remove the entire old ST-426 focal-reference subsection and its figure.
old_84 = next(p for p in doc.paragraphs if p.text.startswith("8.4 ST-426"))
appendix_b = next(p for p in doc.paragraphs if p.text.startswith("附录 B"))
paras = doc.paragraphs
i0 = next(i for i, p in enumerate(paras) if p._p is old_84._p)
i1 = next(i for i, p in enumerate(paras) if p._p is appendix_b._p)
for p in list(paras[i0:i1]):
    remove_paragraph(p)

# Remove old focal references and renumber the references that remain.
for p in list(doc.paragraphs):
    if p.text.startswith("[2] HOLO/OR, Beam Shaper product page") or \
       p.text.startswith("[6] HOLO/OR, TH MATLAB Tutorial"):
        remove_paragraph(p)
for p in doc.paragraphs:
    replacements = {
        "[3] HOLO/OR, FAQ": "[2] HOLO/OR, FAQ",
        "[4] HOLO/OR, Beam Shaper Advantages": "[3] HOLO/OR, Beam Shaper Advantages",
        "[5] HOLO/OR, Top Hat Beam case study": "[4] HOLO/OR, Top Hat Beam case study",
    }
    for old, new in replacements.items():
        if p.text.startswith(old):
            p.text = p.text.replace(old, new, 1)
            for r in p.runs:
                set_run(r, size=9.5)

# Insert the actual MATLAB program process after the existing file-index table.
h8 = next(p for p in doc.paragraphs if p.text.startswith("8. 仿真结果"))
# The source document placed a manual page break immediately after the
# section-7 heading.  Keeping it would leave a nearly blank table-only page.
h7_heading = next(p for p in doc.paragraphs if p.text.startswith("7. MATLAB"))
h7_idx = next(i for i, p in enumerate(doc.paragraphs) if p._p is h7_heading._p)
if h7_idx + 1 < len(doc.paragraphs):
    possible_break = doc.paragraphs[h7_idx + 1]
    if not possible_break.text.strip() and possible_break._p.xpath('.//w:br[@w:type="page"]'):
        remove_paragraph(possible_break)
add_heading(h8, "7.1 主程序入口与调用顺序", 2)
add_body(h8,
    "主程序为 run_HOLOOR_collimated_line_R2019b.m。它先根据自身文件位置解析资料包根目录和结果目录，"
    "创建输出目录，读取统一配置，检查 MATLAB 版本，然后调用准直方案求解器。求解器再依次完成 X 向 ABCD 检查、"
    "厂商出口复场检查、独立双元件替代模型传播和 XZ/YZ 光路绘制。主程序最后写出 MATLAB 版本、时间和状态清单。")
add_code(h8,
    "cfg = common_holoor_collimated_config();\n"
    "result = run_holoor_collimated_solution(cfg, outputDir);\n"
    "% run_holoor_collimated_solution 内部：\n"
    "%   1) X 向 ABCD 与边缘光线\n"
    "%   2) HOLO/OR 出口复场接口（若有）\n"
    "%   3) 独立双元件复场替代模型\n"
    "%   4) XZ/YZ 几何光路图与 CSV/MAT 输出")

add_heading(h8, "7.2 统一参数配置", 2)
add_body(h8,
    "common_holoor_collimated_config.m 是唯一参数源，单位统一为 mm 和 rad。"
    "文件设置 λ=355×10⁻⁶ mm、M²=1.2、激光出口直径 1.0 mm、圆对称 6× 扩束后直径 6.0 mm、"
    "终端 X 宽度 1.5 mm、终端 Y 宽度 1.0×10⁻³ mm、fY=100 mm，并由高斯束腰关系计算 YFL 入射直径 54.240 mm"
    "和含 10% 余量的有效口径 59.664 mm。所有元件 z 位置、焦距和相位验收阈值也在该文件中集中定义。")
add_code(h8,
    "cfg.lambda = 355e-6;       % mm\n"
    "cfg.M2 = 1.2;\n"
    "cfg.moduleOutputXWidth = 6.0;\n"
    "cfg.targetXWidth = 1.5;\n"
    "cfg.targetYWidth = 1.0e-3; % mm = 1.0 um\n"
    "cfg.layout.fXC1 = +160; cfg.layout.fXC2 = -40;\n"
    "cfg.layout.fYFL = +100; cfg.xLastLensPower = 0;")

add_heading(h8, "7.3 X 向 ABCD 与边缘光线验证", 2)
add_body(h8,
    "run_holoor_collimated_solution.m 用 P(d) 和 L(f) 构造 CXC1/CXC2 的无焦矩阵。"
    "对模块出口 ±3 mm、入射角为 0 的两条边缘光线进行传播；结果应在 CXC2 后变为 ±0.75 mm，"
    "且角度仍为 0。YFL 在 X 平面使用单位矩阵 Ix，因此从 z=600 mm 到目标面的传播不引入 X 光焦度。"
    "程序把 C=0、末片 X 光焦度=0、CXC 后宽度、目标宽度和边缘光线角度逐项写入 collimated_x_verification.csv。")
add_code(h8,
    "P = @(d)[1 d; 0 1];\n"
    "L = @(f)[1 0; -1/f 1];\n"
    "Ix = eye(2);\n"
    "MxTelescope = L(-40)*P(120)*L(160);\n"
    "MxToTarget = P(150)*Ix*MxTelescope;\n"
    "rayIn = [-3 3; 0 0];\n"
    "rayAtTarget = MxToTarget*rayIn;")

add_heading(h8, "7.4 HOLO/OR 出口复场数据接口", 2)
add_body(h8,
    "当 01_vendor_model/HOLOOR_CBSM_LINE_355_exit_complex_field.mat 存在时，程序读取 x_mm 和 U_exit；"
    "也兼容 amplitude 与 phase_rad 两数组。有效区按振幅不低于峰值 e⁻¹（等价于强度不低于峰值 e⁻²）选取，"
    "随后展开相位并分别进行二次和一次拟合。一次项换算平均残余倾角；只去除 piston 和一次倾斜后的相位 RMS"
    "保留二次曲率，用于避免把会聚或发散波前误判为准直。")
add_code(h8,
    "active = abs(uVendor) >= exp(-1)*max(abs(uVendor));\n"
    "phi = unwrap(angle(uVendor(active)));\n"
    "fitCoef = polyfit(xAct, phi, 2);\n"
    "linearCoef = polyfit(xAct, phi, 1);\n"
    "phaseRms = rms(phi - polyval(linearCoef,xAct));\n"
    "residualTilt = linearCoef(1)*cfg.lambda/(2*pi);\n"
    "curvature = pi/(cfg.lambda*fitCoef(1));")
add_callout(h8, "验收边界",
    "厂商复场文件缺失时，phaseDataAvailable=0、phaseVerified=0，最终 Vendor acceptance pass 必须保持为 0。"
    "当前接口首先检查模块出口波前；收到正式复场后，还应把该复场通过同一 CXC 传播链复算终端指标，"
    "或以样机终端波前实测关闭最终验收。", "FCE8E6")

add_heading(h8, "7.5 独立双元件相位替代模型", 2)
add_body(h8,
    "simulate_holoor_phase_surrogate.m 用于在厂商相位缺失时进行物理对比，不是厂商代码。"
    "网格为 N=32768、窗口 80 mm、采样间隔约 2.44 µm；输入为 6 mm 全 1/e² 直径的高斯振幅并受 15 mm 口径截断。"
    "程序先用累计能量守恒建立高斯到线形平顶的坐标映射 T(ξ)，积分局部偏转角得到第一相位 φ1，"
    "再计算传播到第二元件前的复场和相位补偿 φ2。")
add_body(h8,
    "当前数值版本把第二元件后的场设置为相位平坦、m=8 的超高斯理想出口场，随后让该复场通过实际后级光路。"
    "因此它验证的是“理想双元件准直出口 + 实际 CXC 后级”是否满足目标，而不是已经完成制造约束优化的 DOE 微结构。"
    "φ1、φ2 是结构原理与后续优化的初值/对照量；若要形成可加工 DOE，仍需迭代传播优化、相位量化、口径、效率和公差分析。")
add_code(h8,
    "xMapped = (outputWidth/2)*erf(sqrt(2)*x/w)/cdfEdge;\n"
    "phi1 = cumtrapz(x, k*sin(atan2(xMapped-x, mapperDistance)));\n"
    "Ebefore2 = propagate(Ein.*exp(1i*phi1), mapperDistance);\n"
    "phi2 = -unwrap(angle(Ebefore2));\n"
    "Eafter2 = exp(-(abs(x)/(outputWidth/2)).^8); % 理想准直出口参考")

add_heading(h8, "7.6 后级复场传播与末片镜子处理", 2)
add_body(h8,
    "替代模型的出口复场先传播到 CXC1，乘以 fX=+160 mm 的薄透镜二次相位；传播 120 mm 后乘以 fX=-40 mm"
    "的薄透镜相位，再传播 150 mm 到目标面。传播函数使用一维角谱传递函数，完整保留振幅和相位。"
    "由于末片 YFL 只对 Y 聚焦，X 传播代码中有意不施加 YFL 的薄透镜相位，这正是保证 X 近似平行的程序实现。")
add_code(h8,
    "E = propagate(Eafter2, zCXC1-zCBSM2);\n"
    "E = E.*exp(-1i*k*x.^2/(2*fXC1));\n"
    "E = propagate(E, zCXC2-zCXC1);\n"
    "E = E.*exp(-1i*k*x.^2/(2*fXC2));\n"
    "Etarget = propagate(E, zTarget-zCXC2);\n"
    "% YFL: X optical power = 0, 因此不乘 X 向透镜相位")

add_heading(h8, "7.7 指标提取、状态和输出", 2)
add_body(h8,
    "wavefrontMetrics 在有效光斑内拟合二次相位，计算平均倾角、局部角度 RMS、去 piston/倾斜后的相位 RMS"
    "和波前曲率半径；thresholdWidth 按强度 e⁻² 阈值提取全宽。独立模型通过条件为：线宽误差≤5%、"
    "平均倾角≤10 µrad、相位 RMS≤0.10 rad、曲率边缘斜率≤50 µrad。geometricPass 表示几何链通过；"
    "designPass 表示几何链与厂商相位或独立模型至少一项通过；acceptancePass 只在几何链和厂商相位均通过时为真。")
add_body(h8,
    "输出包括 collimated_x_abcd.csv、collimated_x_verification.csv、holoor_phase_verification.csv、"
    "holoor_collimated_summary.csv、phase_surrogate_profiles.csv、phase_surrogate_summary.csv、MAT 结果文件及 XZ/YZ/复场图。"
    "plot_holoor_collimated_paths.m 只绘制几何包络，不替代衍射复场判定；无 X 光焦度的 Y 元件在 XZ 图中用虚线表示。")

add_heading(h8, "7.8 MATLAB R2019b 运行方式", 2)
add_body(h8, "在 MATLAB R2019b 命令窗口运行以下命令；主程序会自动更新最终结果目录。")
add_code(h8,
    "cd('D:/work/git_share/laser_sim/deliverables/HOLOOR_CBSM_LINE_355nm_package/02_matlab_code');\n"
    "result = run_HOLOOR_collimated_line_R2019b();\n"
    "result.geometricPass\n"
    "result.designPass\n"
    "result.acceptancePass")

# Add a latest-package appendix before the official references.
appendix_b = next(p for p in doc.paragraphs if p.text.startswith("附录 B"))
add_heading(appendix_b, "附录 A：最终资料包结构", 1)
for text in [
    "01_vendor_model：仅保留 HOLO/OR 出口复场文件格式说明；收到官方复场后放入此目录。",
    "02_matlab_code：仅保留准直线光最终主程序、配置、求解器、独立替代模型和光路绘图代码。",
    "03_simulation_results/collimated_line_R2019b：仅保留最终准直方案结果。",
    "04_official_references：仅保留与准直模块、DOE 原理和采购核对有关的官方资料链接。",
    "HOLOOR_355nm准直线光技术说明.docx：本技术说明。",
]:
    p = insert_before(appendix_b, style="List Bullet")
    set_run(p.add_run(text), size=10)
    p.paragraph_format.space_after = Pt(4)

# The source also carried a manual page break after the result block.  The
# old focal-reference subsection has been removed, so the break would create
# a blank page before the appendices.
for p in list(doc.paragraphs):
    if not p.text.strip() and p._p.xpath('.//w:br[@w:type="page"]'):
        idx = next(i for i, q in enumerate(doc.paragraphs) if q._p is p._p)
        if idx + 1 < len(doc.paragraphs) and doc.paragraphs[idx + 1].text.startswith("附录 A"):
            remove_paragraph(p)

doc.core_properties.title = "355 nm HOLO/OR 准直线光技术说明"
doc.core_properties.subject = "最终准直线光方案、DOE 原理、数学模型与 MATLAB R2019b 程序过程"
doc.core_properties.comments = "已移除旧焦面方案及旧 HOLO/OR 厂商程序内容"
doc.save(out)
print(out)
