import os
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


path = os.environ["DOCX_PATH"]
doc = Document(path)

title = doc.paragraphs[0]
title.text = "355 nm HOLO/OR 准直线光技术说明"
for run in title.runs:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(24)
    run.bold = True

for p in doc.paragraphs:
    if "phaseRms = rms(phi - polyval(linearCoef,xAct));" in p.text:
        for run in p.runs:
            run.text = run.text.replace(
                "phaseRms = rms(phi - polyval(linearCoef,xAct));",
                "phaseRms = sqrt(mean((phi-polyval(linearCoef,xAct)).^2));")

doc.save(path)
print(path)
